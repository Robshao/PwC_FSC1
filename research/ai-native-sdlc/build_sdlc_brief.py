# AI 原生軟體開發流程：給 Sam 的三頁精簡版（白話、完整句子，可直接對客戶說明怎麼落地）
# 內容整理自 Anthropic《The AI-Native SDLC Playbook》；完整版見 build_sdlc_deck.py
# 用法：python3 build_sdlc_brief.py <範本.pptx> <輸出.pptx>
import sys, copy, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TPL, OUT = sys.argv[1], sys.argv[2]
OR = RGBColor(0xFD, 0x51, 0x08); OR2 = RGBColor(0xFE, 0x7C, 0x39); TINT = RGBColor(0xFF, 0xE8, 0xD4)
# 範本佈景主題色：tx1 黑、lt2 淺灰、accent6 線條灰、accent4 中灰
INK = RGBColor(0x00, 0x00, 0x00); GREY = RGBColor(0x59, 0x59, 0x59); LG = RGBColor(0xEB, 0xEB, 0xEB); LINE = RGBColor(0xCB, 0xD1, 0xD6)
MID = RGBColor(0xA1, 0xA8, 0xB3); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# 範本字型：標題 Georgia、內文 Arial、中文一律微軟正黑體
FONT = 'Arial'; HEAD_FONT = 'Georgia'; EA_FONT = '微軟正黑體'
SYMBOLS = set('●◐○✓✗→①②③④⑤⑥⑦⚠↑σ')
FOOT = '對外說明版｜內容整理自 Anthropic《The AI-Native SDLC Playbook》｜時間縮短幅度為該手冊所述預期，實際以試點量測為準'
SRC = 'Anthropic《The AI-Native SDLC Playbook》'

def style_font(r, head=False):
    # Arial 沒有 ●◐○ 等符號，這類文字改用中文字型顯示
    r.font.name = EA_FONT if SYMBOLS & set(r.text) else (HEAD_FONT if head else FONT)
    rPr = r._r.get_or_add_rPr()
    for e in rPr.findall(qn('a:ea')): rPr.remove(e)
    rPr.find(qn('a:latin')).addnext(rPr.makeelement(qn('a:ea'), {'typeface': EA_FONT}))

p = Presentation(TPL)
lst = p.slides._sldIdLst
for sid in list(lst):
    p.part.drop_rel(sid.rId); lst.remove(sid)
L = {l.name: l for l in p.slide_masters[2].slide_layouts}
# 移除範本遺留的投影片章節（PowerPoint 左側的章節名稱）
for ext in list(p.part._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}ext')):
    if ext.find('{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst') is not None:
        ext.getparent().remove(ext)

def set_ph(s, idx, text, size=None, bold=None):
    for ph in s.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            for para in ph.text_frame.paragraphs:
                for r in para.runs:
                    style_font(r, head=idx == 0)
                    if size: r.font.size = Pt(size)
                    if bold is not None: r.font.bold = bold
            return ph

def tb(s, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, margin=0.06, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'): setattr(tf, f'margin_{side}', Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg in segs:
            t, st = (seg, {}) if isinstance(seg, str) else seg
            r = para.add_run(); r.text = t
            style_font(r); f = r.font; f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color)
            if 'link' in st:
                r.hyperlink.address = st['link']; f.underline = True
        para.space_after = Pt(2)
    return sh

def circle(s, x, y, d, text, fill=OR, size=12):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    r = para.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; style_font(r)

def hline(s, x1, y1, x2, y2, color=LINE, width=1.0):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width)

def arrow(s, x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, fill=OR):
    a = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = fill; a.line.fill.background(); a.shadow.inherit = False

def footer(s):
    # python-pptx 新增投影片不會帶入頁尾與頁碼，從版面複製（頁碼為自動欄位）
    have = {ph.placeholder_format.idx for ph in s.placeholders}
    for ph in s.slide_layout.placeholders:
        if ph.placeholder_format.idx in (3, 4) and ph.placeholder_format.idx not in have:
            s.shapes._spTree.append(copy.deepcopy(ph._element))
    set_ph(s, 3, FOOT)

def content(title):
    s = p.slides.add_slide(L['只有標題'])
    set_ph(s, 0, title, 20, bold=True); footer(s)
    return s

def note(s, text, y=6.66):
    tb(s, 0.45, y, 12.45, 0.32, text, size=9, color=GREY, margin=0.03)

def grid(s, x, y, cw, head, rows, rh, size=10.5, first_bold=True, head_h=0.4, hot=()):
    # 以文字框排成的表格：表頭橘底白字、隔列淺灰；hot 為要強調的列
    xx = x
    for j, h in enumerate(head):
        tb(s, xx, y, cw[j] - 0.04, head_h, h, size=size, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
        xx += cw[j]
    for i, row in enumerate(rows):
        yy = y + head_h + 0.04 + i * rh; xx = x
        bg = TINT if i in hot else (LG if i % 2 == 0 else WHITE)
        for j, v in enumerate(row):
            first = j == 0 and first_bold
            tb(s, xx, yy, cw[j] - 0.04, rh - 0.04, v, size=size, bold=first, color=OR if first else INK,
               fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
            xx += cw[j]

def say(s, x, y, w, h, text, size=12):
    # 「Sam 可以這樣說」引言框
    tb(s, x, y, 0.1, h, '', fill=OR)
    tb(s, x + 0.1, y, w - 0.1, h, [[('Sam 可以這樣說', {'bold': True, 'color': OR, 'size': 11})], text], size=size, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.09)

# ───────── 第 1 頁 一頁說清楚 ─────────
s = content('一頁說清楚：AI 讓寫程式變快，我們協助客戶把整個開發流程一起改，關鍵決定仍由客戶的人來做')
BLOCKS = [('這是什麼', '軟體從想法到上線，通常要經過規劃、設計、開發、測試、上線、維運六個階段。AI 原生的做法，是讓 AI 參與每一個階段：整理需求、起草規格、寫程式、跑測試、做第一輪審查、盯系統異常。每個階段結束都留下一份檔案，下一個階段直接接手，不必再開會交接。'),
          ('為什麼現在要做', '現在 AI 寫程式已經很快，但多數公司的需求討論、程式審查、上線核准還是照舊，整體交付時間沒有跟著縮短；程式量變多之後，審查反而排得更長。要真正變快，流程本身也要一起改。')]
y = 1.45
for h, t in BLOCKS:
    tb(s, 0.45, y, 6.4, 0.38, h, size=13, bold=True, color=OR)
    box = tb(s, 0.45, y + 0.4, 6.4, 1.55 if h == '這是什麼' else 1.2, t, size=12, fill=LG, margin=0.14)
    y += 2.1 if h == '這是什麼' else 1.75
tb(s, 0.45, 5.3, 6.4, 0.38, '客戶得到什麼', size=13, bold=True, color=OR)
GAIN = [('更快', '手冊預期從想法到需求定稿可由數週縮短到數小時，實際以試點量測'),
        ('查得到', '誰提出、AI 做了什麼、誰核准，全部有紀錄'),
        ('守得住', '公司規範寫成 AI 會照做的規則，上線一定要人授權')]
for i, (h, t) in enumerate(GAIN):
    x = 0.45 + i * 2.15
    tb(s, x, 5.72, 2.08, 0.9, [[(h, {'bold': True, 'color': OR, 'size': 12})], t], size=10, fill=LG, margin=0.09)
say(s, 7.15, 1.45, 5.75, 3.65,
    '「各位現在用 AI 寫程式應該已經快很多，但可能也發現，需求討論、程式審查、上線核准還是跟以前一樣慢，整體並沒有快多少。'
    '我們要談的，是把整個開發流程一起重新設計：AI 參與每一個步驟，每一步都留下紀錄；像『這個需求要不要做』、『這次改動能不能上線』這類關鍵決定，仍然由你們指定的人來做。」', size=12.5)
tb(s, 7.15, 5.3, 5.75, 1.32, [[('一句話收尾：', {'bold': True, 'color': OR})],
                              '「AI 幫你把每一步做得更快，每一步也都留下紀錄；要不要上線，永遠由你的人決定。」'],
   size=12.5, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.16)
note(s, f'來源：{SRC}，第 3–7 頁。「數週到數小時」為手冊所述預期，非實測。')

# ───────── 第 2 頁 實際怎麼運作 ─────────
s = content('實際怎麼運作：以「理賠進度自助查詢」為例，人負責說清楚問題、決定要不要做、決定能不能上線')
STEPS = [['① 提出問題', '理賠部門同仁直接跟 AI 說明：客戶常打電話問理賠進度，客服約三分之一通話時間都在回答這件事。AI 會像分析師一樣追問範圍、使用者和限制，整理成需求說明。', '提出的同仁確認內容', '需求說明'],
         ['② 決定要不要做', '產品負責人看過 AI 整理的需求說明，修正誤解後，決定接受或退回。', '★ 產品負責人', '接受紀錄'],
         ['③ 起草規格', 'AI 依需求說明寫出規格，自動套用公司的資安、法遵、品牌規範，例如「入口網站不能新增個資」。規範互相衝突的地方，AI 會標出來交給規範負責人。', '產品負責人審核', '規格書'],
         ['④ 寫實作計畫', 'AI 先寫出要改哪些地方、先後順序、可能的風險（例如核心系統有流量上限，要加快取）以及怎麼證明完成。工程師修正、核准之前，AI 不能改任何程式。', '工程師（高風險找架構師）', '實作計畫'],
         ['⑤ 寫程式與測試', 'AI 照計畫寫程式並自己跑測試，確認四種理賠狀態都正確、畫面和設計稿一致，全部通過才交給人。', 'AI 自行驗證', '程式與測試結果'],
         ['⑥ 審查與上線', 'AI 先依公司審查規則檢查錯誤、資安與合規問題並排出嚴重程度；程式負責人判斷是否符合原意、風險能否接受；最後由發布經理授權上線。', '★ 程式負責人、發布經理', '審查與核准紀錄'],
         ['⑦ 上線後監看', '錯誤率超出正常範圍時，AI 先診斷原因，必要時執行事先演練過的退回上一版，並把發現寫成新的需求說明，重新進入流程。', '服務負責人', '事件紀錄']]
grid(s, 0.45, 1.4, [1.75, 7.05, 2.1, 1.55], ['步驟', '發生什麼（白話）', '誰決定', '留下的紀錄'], STEPS, 0.6, size=10.5, head_h=0.38, hot=(1, 5))
say(s, 0.45, 6.06, 12.45, 0.6, '「整個過程裡，起草、計畫、寫程式、測試、初審都由 AI 先做；你們的人只在關鍵點做決定，而且每一步都有紀錄可查。」', size=11.5)
note(s, f'★ 與淺橘底為人做決定的關卡。數字（三分之一通話時間）為手冊示例，非客戶資料。來源：{SRC}，第 12–20、37–52 頁。', y=6.7)

# ───────── 第 3 頁 怎麼導入 ─────────
s = content('我們怎麼帶客戶導入：先挑一個小試點跑一次，量到成效再擴大；客戶常問的問題這樣回答')
tb(s, 0.45, 1.4, 6.3, 0.42, '導入四步驟', size=13, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
ROLL = [('先挑一個試點', '選一個規格清楚、影響範圍小、已經有測試的日常需求，降低第一次的風險。'),
        ('把規矩寫下來', '準備需求說明範本與團隊說明檔，把一條常被漏掉的公司規範寫成 AI 會照做的規則，並用自動閘門守住現有的核准關卡。'),
        ('帶著跑一次', '試點需求從頭到尾走一遍上一頁的七個步驟。先由人一步步帶著 AI 做，確認每個關卡都有人核准、都有紀錄。'),
        ('量測再擴大', '試點前先記下現況：需求定稿時間、審查等待時間、上線失敗率。試點後用同一套指標比較，數字站得住再擴大，最後才讓流程自動串接。')]
for i, (h, t) in enumerate(ROLL):
    y = 1.9 + i * 1.0
    circle(s, 0.45, y + 0.2, 0.5, str(i + 1), size=13)
    tb(s, 1.05, y, 5.7, 0.92, [[(h, {'bold': True, 'size': 12})], t], size=10.5, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
tb(s, 7.05, 1.4, 5.85, 0.42, '客戶常問的四個問題', size=13, bold=True, color=WHITE, fill=INK, anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
QA = [('會不會失控？', 'AI 只能走核准過的路徑，正式上線一定要指定的人授權，每一次放行或阻擋都有時間紀錄。'),
      ('出事誰負責？', '需要判斷的決定都由人做並留下紀錄；寫程式的 AI 不能核准自己的程式，職責分離不變。'),
      ('要換掉現有的 Jira 嗎？', '不用。每種紀錄指定一個「以誰為準」的系統，其他系統放連結即可。'),
      ('機密資料安全嗎？', 'IT 可統一禁止 AI 讀取金鑰與機密檔、限制對外連線、在隔離環境執行，也可走公司既有的雲端合約。')]
for i, (q, a) in enumerate(QA):
    y = 1.9 + i * 1.0
    tb(s, 7.05, y, 5.85, 0.92, [[(q, {'bold': True, 'size': 12, 'color': OR})], a], size=10.5, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
say(s, 0.45, 5.98, 12.45, 0.68, '「我們建議先挑一個小範圍的需求，用你們現有的版本紀錄和 CI 數據，量出導入前後的差別；數字站得住，再擴大到更多團隊。」', size=11.5)
note(s, '推案前：客戶若同時是所內審計客戶，先與 Risk & Quality 確認獨立性。完整說明、法規、服務模組與術語對照見 25 頁完整版。', y=6.7)

p.save(OUT)
print(OUT, len(p.slides), 'slides')
