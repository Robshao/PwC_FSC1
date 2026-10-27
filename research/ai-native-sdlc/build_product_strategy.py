# AI 原生開發：從流程到產品（內部會議，22 頁：含議程、五個章節轉場頁與附錄轉場頁）
# 結論先行：差距分析（AI 原生 SDLC 對照現行流程）→ 優先順序（多常發生 × 需要判斷）→ 市場驗證（同業產品化）→ 建議 → 各方意義 → 決策
# 不提及任何個人姓名；同業公司名稱只放附錄
# 用法：python3 build_product_strategy.py <範本.pptx> <輸出.pptx>
import sys, copy, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TPL, OUT = sys.argv[1], sys.argv[2]
OR = RGBColor(0xFD, 0x51, 0x08); OR2 = RGBColor(0xFE, 0x7C, 0x39); TINT = RGBColor(0xFF, 0xE8, 0xD4)
# 範本佈景主題色：tx1 黑、lt2 淺灰、accent6 線條灰、accent4 中灰
INK = RGBColor(0x00, 0x00, 0x00); GREY = RGBColor(0x59, 0x59, 0x59); LG = RGBColor(0xEB, 0xEB, 0xEB); LINE = RGBColor(0xCB, 0xD1, 0xD6)
MID = RGBColor(0xA1, 0xA8, 0xB3); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# 範本字型：標題 Georgia、內文 Arial、中文一律微軟正黑體
FONT = 'Arial'; HEAD_FONT = 'Georgia'; EA_FONT = '微軟正黑體'
SYMBOLS = set('●◐○✓✗→↑①②③④⑤⑥⑦⚠')
FOOT = '內部會議｜AI 原生開發：從流程到產品｜AI 協助產出，需人工確認後再對外'
V3 = '現行開發流程 v3'
PB = 'Anthropic《The AI-Native SDLC Playbook》'
SMR = 'MIT Sloan Management Review〈How to Turn Professional Services Into Products〉，2024'

def style_font(r, head=False):
    # Arial 沒有 ●◐○ 等符號，這類文字改用中文字型顯示
    r.font.name = EA_FONT if SYMBOLS & set(r.text) else (HEAD_FONT if head else FONT)
    rPr = r._r.get_or_add_rPr()
    for e in rPr.findall(qn('a:ea')): rPr.remove(e)
    rPr.find(qn('a:latin')).addnext(rPr.makeelement(qn('a:ea'), {'typeface': EA_FONT}))

p = Presentation(TPL)
lst = p.slides._sldIdLst
for sid in list(lst):
    p.part.drop_rel(sid.rId); lst.remove(sid)
L = {l.name: l for l in p.slide_masters[2].slide_layouts}
# 移除範本遺留的投影片章節（PowerPoint 左側的章節名稱）
for ext in list(p.part._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}ext')):
    if ext.find('{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst') is not None:
        ext.getparent().remove(ext)

def set_ph(s, idx, text, size=None, bold=None):
    for ph in s.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            for para in ph.text_frame.paragraphs:
                for r in para.runs:
                    style_font(r, head=idx == 0)
                    if size: r.font.size = Pt(size)
                    if bold is not None: r.font.bold = bold
            return ph

def tb(s, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, margin=0.06, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'): setattr(tf, f'margin_{side}', Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg in segs:
            t, st = (seg, {}) if isinstance(seg, str) else seg
            r = para.add_run(); r.text = t
            style_font(r); f = r.font; f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color)
            if 'link' in st:
                r.hyperlink.address = st['link']; f.underline = True
        para.space_after = Pt(2)
    return sh

def circle(s, x, y, d, text, fill=OR, size=12):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    r = para.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; style_font(r)

def hline(s, x1, y1, x2, y2, color=LINE, width=1.0):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width)

def arrow(s, x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, fill=OR):
    a = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = fill; a.line.fill.background(); a.shadow.inherit = False

def footer(s):
    # python-pptx 新增投影片不會帶入頁尾與頁碼，從版面複製（頁碼為自動欄位）
    have = {ph.placeholder_format.idx for ph in s.placeholders}
    for ph in s.slide_layout.placeholders:
        if ph.placeholder_format.idx in (3, 4) and ph.placeholder_format.idx not in have:
            s.shapes._spTree.append(copy.deepcopy(ph._element))
    set_ph(s, 3, FOOT)

def content(title):
    s = p.slides.add_slide(L['只有標題'])
    set_ph(s, 0, title, 20, bold=True); footer(s)
    return s

def note(s, text, y=6.66):
    tb(s, 0.45, y, 12.45, 0.32, text, size=9, color=GREY, margin=0.03)

def section(num, title, ask, see):
    # 章節轉場頁：沿用範本「章節標題」版面，左上方說明這一章要回答什麼、會看到什麼
    s = p.slides.add_slide(L['章節標題'])
    set_ph(s, 0, title, 30); set_ph(s, 1, num, 180); footer(s)
    tb(s, 0.45, 0.7, 7.9, 2.75, [[('這一章要回答', {'bold': True, 'color': OR, 'size': 13})], [(ask, {'size': 17, 'bold': True})],
                                 [('', {'size': 6})], [('接下來你會看到', {'bold': True, 'color': OR, 'size': 13})]] + [[('・' + x, {'size': 14})] for x in see],
       fill=WHITE, margin=0.22)
    return s

def grid(s, x, y, cw, head, rows, rh, size=10.5, first_bold=True, head_h=0.4, hot=()):
    # 以文字框排成的表格：表頭橘底白字、隔列淺灰；hot 為要強調的列
    xx = x
    for j, h in enumerate(head):
        tb(s, xx, y, cw[j] - 0.04, head_h, h, size=size, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
        xx += cw[j]
    for i, row in enumerate(rows):
        yy = y + head_h + 0.04 + i * rh; xx = x
        bg = TINT if i in hot else (LG if i % 2 == 0 else WHITE)
        for j, v in enumerate(row):
            first = j == 0 and first_bold
            tb(s, xx, yy, cw[j] - 0.04, rh - 0.04, v, size=size, bold=first, color=OR if first else INK,
               fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
            xx += cw[j]

def callout(s, y, label, text, h=0.62, size=11.5):
    tb(s, 0.45, y, 0.1, h, '', fill=OR)
    tb(s, 0.55, y, 12.35, h, [[(label, {'bold': True, 'color': OR}), (text, {})]], size=size, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.16)


def draw_loop(s, cx, cy, R, d=1.05):
    # 自繪「循環」圖：六個階段排成一圈，中間是 AI 協作；不使用任何廠商的圖
    names = ['需求分析', 'SA&D', '開發', '測試', '上版／交付', '維運']
    pts = [(cx + R * math.cos(math.radians(-90 + 60 * i)), cy + R * math.sin(math.radians(-90 + 60 * i))) for i in range(6)]
    for i in range(6):
        (x1, y1), (x2, y2) = pts[i], pts[(i + 1) % 6]
        hline(s, x1, y1, x2, y2, OR2, 2.5)
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.75), Inches(cy - 0.75), Inches(1.5), Inches(1.5))
    c.fill.solid(); c.fill.fore_color.rgb = OR; c.line.fill.background(); c.shadow.inherit = False
    tb(s, cx - 0.75, cy - 0.4, 1.5, 0.8, ['AI 協作', [('參與每一步', {'size': 9.5})]], size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
    for (x, y), n in zip(pts, names):
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d))
        o.fill.solid(); o.fill.fore_color.rgb = WHITE if n != '維運' else TINT; o.line.color.rgb = OR; o.line.width = Pt(1.5); o.shadow.inherit = False
        tb(s, x - d / 2, y - 0.25, d, 0.5, n, size=10.5 if len(n) < 5 else 9.5, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)

CHECK = [['是否屬作業委外、是否需依規定辦理', '保險業作業委託他人處理內部作業制度及程序辦法（2025-12-31 施行）', '導入前判斷並留下紀錄；AI 工具納入委外管理', '診斷'],
         ['雲端資料存放、雲端業者管理與查核', '保險業作業委外使用雲端服務自律規範', '使用公司核准的雲端合約；保留查核所需紀錄', '診斷'],
         ['AI 工具登錄 AI 清冊並做風險評估', '保險業運用人工智慧系統自律規範（2025-04-10 備查）', '每個 AI 工具與用途登錄；高風險用途加審', '診斷'],
         ['關鍵決定由人做，保險公司負最終責任', '同上；金管會 AI 指引「治理及問責」', '五個人工關卡，各有指定核准人', '①–⑤'],
         ['個資與客戶資料不進入 AI 可讀範圍', '個人資料保護法；金管會 AI 指引「隱私及客戶權益」', '資料分級；集中設定禁止讀取機密檔與金鑰', '規格、開發'],
         ['程式經資安審查，出事可退回上一版', '金管會 AI 指引「系統穩健性與安全性」', 'AI 資安初審＋人工審查；退版流程事先演練', '④ 合併、⑤ 上線'],
         ['誰要求、AI 產出什麼、誰核准，查得到', '金管會 AI 指引「透明性與可解釋性」', '需求說明、規格、計畫與核准紀錄版本化保存', '全程']]

def checklist_page(title, foot):
    s = content(title)
    grid(s, 0.45, 1.45, [3.2, 3.85, 3.85, 1.55], ['檢查項目', '依據', '開發流程中怎麼落實', '在哪一步檢查'], CHECK, 0.62, size=10.5)
    note(s, foot, y=6.42)
    return s

KPI = [['需求定稿時間', '從第一次討論到需求說明定稿', '試點前量測', '雙方約定'],
       ['第一次審查等待時間', '從提交程式到收到第一次審查意見', '試點前量測', '雙方約定'],
       ['有完整核准紀錄的改動比例', '比對版本紀錄與核准紀錄', '試點前量測', '雙方約定']]

def criteria_page(title):
    s = content(title)
    grid(s, 0.45, 1.45, [2.8, 3.6, 1.6, 1.6], ['指標（可依痛點替換）', '怎麼量', '基準', '合格線'], KPI, 0.75, size=11.5)
    tb(s, 10.25, 1.45, 2.65, 2.7, [[('結果怎麼用', {'bold': True, 'color': OR, 'size': 12.5})],
                                   [('達標：', {'bold': True}), ('擴大到下一個團隊。', {})],
                                   [('未達標：', {'bold': True}), ('一起檢討原因，調整後再試或停止，不擴大。', {})]], size=11.5, fill=TINT, margin=0.14)
    tb(s, 0.45, 4.45, 12.45, 1.2, [[('為什麼這樣做', {'bold': True, 'color': OR, 'size': 12.5})],
                                   '效益不先寫數字，而是試點開始前一起訂好指標和合格線；數據大多已在版本紀錄與 CI 系統中，不需要另外填表。試點結束，用同一套指標比較前後。'],
       size=12, fill=LG, margin=0.16)
    return s

def section_tag(s, text, i):
    tb(s, 10.4, 1.0, 2.5, 0.3, text, size=9.5, bold=True, color=WHITE, fill=[OR, OR2, INK, MID, GREY][i], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)

# ───────── 1 封面 ─────────
s = p.slides.add_slide(L['標題投影片'])
set_ph(s, 0, 'AI 原生開發：\n從流程到產品', 28)
set_ph(s, 1, 'FSC1 內部會議｜2026 年 10 月', 12)

# ───────── 2 一頁看懂 ─────────
s = content('一頁看懂：補強現行流程、先產品化「常發生、判斷少」的工作，對外先推診斷＋試點')
SCQ = [('情境', 'AI 讓寫程式變快，但需求、審查、上線還是人工速度；同業已經把類似能力做成產品在賣。'),
       ('問題', '我們已有一套 AI 協作開發流程，要怎麼把它變成可以重複銷售、重複交付的服務？'),
       ('建議', '補強流程的三個缺口，從四象限挑出最適合產品化的工作，以在地法規與可稽核關卡切入市場，先推「診斷＋陪你做試點」。')]
for i, (h, t) in enumerate(SCQ):
    y = 1.45 + i * 1.12
    tb(s, 0.45, y, 1.3, 1.0, h, size=14, bold=True, color=WHITE, fill=OR if i == 2 else INK, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    tb(s, 1.8, y, 5.0, 1.0, t, size=11.5, fill=TINT if i == 2 else LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
PIL3 = [('差距分析', '現行流程已涵蓋主要工作；補強關卡核准人、規則自動執行、上線後回饋', '第 5–7 頁'),
        ('優先順序', '「常發生、判斷少」的工作最適合先交給 AI、先做成工具包', '第 9–10 頁'),
        ('市場驗證', '同業主流是「預建資產＋顧問導入」；在地法規與可稽核關卡是我們的切入點', '第 12–13 頁')]
tb(s, 7.05, 1.45, 5.85, 0.42, '三個支持理由', size=12.5, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
for i, (h, t, pg) in enumerate(PIL3):
    tb(s, 7.05, 1.93 + i * 0.95, 5.85, 0.88, [[(h + '　', {'bold': True, 'color': OR}), (t, {}), ('　' + pg, {'color': GREY, 'size': 9.5})]], size=11, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
callout(s, 5.0, '今天要決定：', '四象限分類、試點案例、對外先推的方案、價格區間與人力、法規檢查清單的審閱（第 19 頁）。', h=0.7, size=12)
note(s, '本簡報不涉及任何客戶資料；AI 協助產出的內容對外前一律人工確認。', y=5.85)

# ───────── 議程 ─────────
s = content('議程')
AG = [('', '一頁看懂', '情境、問題、建議與今天要決定的事', '2'),
      ('01', '差距分析', '理想的 AI 原生流程對照現行流程：保留什麼、補強什麼', '4–7'),
      ('02', '優先順序', '用「多常發生 × 需要判斷」決定哪些工作先交給 AI、先做成產品', '8–10'),
      ('03', '市場驗證', '同業怎麼把保險顧問服務做成產品，我們從哪裡切入', '11–13'),
      ('04', '建議', '入門診斷加三種合作方式、試點成功標準', '14–16'),
      ('05', '下一步', '對各團隊的意義、今天要拍板的事', '17–19'),
      ('附錄', '同業來源、法規檢查清單', '', '20–22')]
for i, (n, h, d, pg) in enumerate(AG):
    y = 1.4 + i * 0.7
    if n: circle(s, 0.6, y + 0.08, 0.5, n, fill=OR if n != '附錄' else MID, size=12 if n != '附錄' else 10)
    tb(s, 1.35, y, 3.0, 0.62, h, size=15, bold=True, color=OR if n == '05' else INK, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 4.4, y, 7.2, 0.62, d, size=11.5, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 11.65, y, 1.25, 0.62, pg, size=15, bold=True, color=OR, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    hline(s, 1.35, y + 0.66, 12.9, y + 0.66)
note(s, '右側數字為頁碼。每一章開頭有轉場頁，說明這一章要回答的問題。', y=6.42)

# ───────── 章節 01 ─────────
section('01', '差距分析', '理想的 AI 原生開發流程長什麼樣？現行流程離它有多遠？', ['從直線變循環的方向', '六個階段：保留什麼、補強什麼', '五個 Review Gate 補上核准人與紀錄'])

# ───────── 3 方向：直線變循環 ─────────
s = content('方向：開發流程要從「直線交棒」變成「循環」，AI 參與每一步，人在關卡上做決定')
section_tag(s, '差距分析', 0)
tb(s, 0.45, 1.45, 2.1, 0.4, '以前：一條直線', size=12, bold=True, color=GREY, margin=0.02)
for i, n in enumerate(['需求分析', 'SA&D', '開發', '測試', '上版／交付', '維運']):
    tb(s, 0.45, 1.9 + i * 0.66, 2.1, 0.5, n, size=11, color=WHITE, fill=MID, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
tb(s, 0.45, 5.9, 2.1, 0.5, '做完一段才交棒，退回一次就要重來', size=9.5, color=GREY, margin=0.02)
tb(s, 3.1, 1.45, 4.5, 0.4, '現在：一個循環', size=12, bold=True, color=OR, margin=0.02)
draw_loop(s, 5.35, 3.85, 1.75)
WHY = [('瓶頸移位', '寫程式變快後，需求、審查、上線這些人工步驟成了瓶頸。'),
       ('交接靠檔案', '每一步交出一份檔案，下一步直接接手，不必開會交接。'),
       ('人在關卡上', 'AI 起草、檢查、執行；接受需求、核准上線由指定的人做。')]
for i, (h, t) in enumerate(WHY):
    tb(s, 8.0, 1.45 + i * 1.45, 4.9, 1.3, [[(h, {'bold': True, 'color': OR, 'size': 13})], t], size=11.5, fill=LG if i < 2 else TINT, margin=0.14)
note(s, f'概念依 {PB}，第 3–9 頁；圖為自繪，不綁特定 AI 工具。', y=6.3)

# ───────── 4 差距：六階段對照 ─────────
s = content('差距：現行流程已涵蓋主要工作，這次補強三件事：關卡有人簽、規則自動執行、上線後回饋')
section_tag(s, '差距分析', 0)
ROWS = [['需求分析', 'AI 產出 Prototype、欄位清單、流程圖，IT Review', '需求說明由 BA 主導對焦；「接受需求」指定核准人'],
        ['SA&D', 'AI 產出 Use Case、ERD、API 規格、Test Plan', '資安、法遵規範寫成 AI 自動套用的規則'],
        ['開發', '先 plan 再 code、限制修改範圍、Code Review', '計畫核准留紀錄；團隊慣例寫成 AI 每次都讀的說明檔'],
        ['測試', '正向、反向、邊界案例；案例綁需求編號', '自動測試通過才能進入審查'],
        ['上版／交付', 'Release Note、Deploy Guide、rollback 計畫', '合併與上線由自動關卡把關，沒有核准做不到'],
        ['維運', '—（流程到上版為止）', '上線後異常由 AI 先診斷，寫成新需求回到第一步']]
grid(s, 0.45, 1.45, [1.6, 5.2, 5.65], ['階段', '現行流程已有（保留）', '這次補強'], ROWS, 0.62, size=11, hot=(5,))
callout(s, 5.75, '一句話：', '不是重做，而是在現行流程上加三樣東西，讓它從「團隊的做法」變成「可以交給客戶的方法」。', h=0.6, size=12)
note(s, f'對照基準：{V3}與 {PB}。', y=6.42)

# ───────── 5 補強：關卡 ─────────
s = content('補強一：現行五個 Review Gate 全部保留，每一關補上「誰核准」和「怎麼確保一定經過」')
section_tag(s, '差距分析', 0)
ROWS = [['需求文件確認', '① 接受需求', '產品負責人', '接受或退回都留紀錄'],
        ['SA／SD 文件確認', '② 確認規格', '產品負責人（高風險加技術主管）', '規格與需求說明一起存檔'],
        ['程式與相關文件完成', '③ 核准計畫', '工程師（高風險找架構師）', '計畫核准前 AI 不能修改程式'],
        ['完整測試報告與紀錄', '④ 核准合併', '程式負責人', '沒有核准不能合併'],
        ['上版／交付', '⑤ 授權上線', '發布經理', '沒有授權，上線指令會被擋下']]
grid(s, 0.45, 1.45, [2.7, 2.0, 3.6, 4.15], ['現行 Gate', '對應關卡', '誰核准', '怎麼確保一定經過'], ROWS, 0.68, size=11.5)
tb(s, 0.45, 5.4, 12.45, 0.85, [[('Review Checklist 同步補四項：', {'bold': True, 'color': OR}), ('法遵（個資、資料範圍）、職責分離（寫程式的 AI 不能自己核准）、測試證據、紀錄完整（誰核准、何時、套用哪版規範）。', {})]], size=11.5, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.16)
note(s, f'關卡設計依 {PB}，第 13、17、20、39、46 頁。', y=6.4)

# ───────── 章節 02 ─────────
section('02', '優先順序', '流程裡哪些工作最適合先交給 AI、先做成產品？', ['「多常發生 × 需要判斷」四象限分類', '每個象限對應的做法與產品元件'])

# ───────── 6 四象限 ─────────
s = content('優先順序：用「多常發生 × 需要判斷」把現行流程的 AI 任務分成四類')
section_tag(s, '優先順序', 1)
MX, MY, QW, QH = 1.15, 1.5, 3.6, 2.15
Q = [(0, 0, '常發生 × 判斷少', 'AI 先做，IT 確認', '欄位清單、Test Case、Test Data、Unit Test、註解與文件、Release Note', TINT),
     (1, 0, '常發生 × 判斷多', 'AI 起草，IT 逐項審查', '程式碼、API 規格、ERD、Code Review 風險檢查、測試失敗原因分析', LG),
     (0, 1, '少發生 × 判斷少', '做成範本，重複使用', 'Deploy Guide、User Guide、Migration Script、部署腳本', LG),
     (1, 1, '少發生 × 判斷多', '人決定，AI 只提供資料', '接受需求、架構選型、核准合併、授權上線、事件處置', WHITE)]
for c, r, h, rule, items, bg in Q:
    x, y = MX + c * (QW + 0.06), MY + r * (QH + 0.06)
    tb(s, x, y, QW, QH, '', fill=bg, line=LINE)
    tb(s, x + 0.12, y + 0.08, QW - 0.24, 0.62, [[(h, {'bold': True, 'size': 11, 'color': GREY})], [(rule, {'bold': True, 'size': 13, 'color': OR})]], margin=0.02)
    tb(s, x + 0.12, y + 0.76, QW - 0.24, QH - 0.84, items, size=11.5, margin=0.02)
tb(s, 0.45, MY, 0.65, QH, '常發生', size=11, bold=True, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
tb(s, 0.45, MY + QH + 0.06, 0.65, QH, '少發生', size=11, bold=True, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
tb(s, MX, MY + 2 * QH + 0.1, 2 * QW + 0.06, 0.3, '需要判斷：少 ─────────────→ 多', size=10.5, bold=True, color=GREY, align=PP_ALIGN.CENTER, margin=0.0)
tb(s, 8.75, 1.5, 4.15, 4.4, [[('怎麼分', {'bold': True, 'color': OR, 'size': 13})],
                             [('多常發生：', {'bold': True}), ('每個專案、每次改動都會做嗎？越常發生，AI 越能學會、省下的工越多。', {})],
                             [('需要判斷：', {'bold': True}), ('要多少專業判斷？做錯的代價多大？判斷越多，越要留給人。', {})],
                             [('', {'size': 6})],
                             [('請 IT 團隊現場確認：', {'bold': True, 'color': OR}), ('分類是依現行流程各階段「AI 可以做什麼」整理的初步判斷，有不同意見今天就調整。', {})]],
   size=11.5, fill=LG, margin=0.16)
note(s, f'方法依 {SMR}：把服務拆成活動，依頻率與認知複雜度評估產品化潛力。', y=6.45)

# ───────── 7 象限 → 產品 ─────────
s = content('每個象限的做法不同，也決定了哪些能做成產品：左上角最先、右下角永遠留給人')
section_tag(s, '優先順序', 1)
ROWS = [['常發生 × 判斷少', 'AI 先做，IT 確認', '高：最先做', '工具包：範本＋規範規則，自動產出'],
        ['常發生 × 判斷多', 'AI 起草，IT 逐項審查', '中：方法可產品化，判斷不行', '陪你做：我們帶客戶建立審查規則'],
        ['少發生 × 判斷少', '做成範本，重複使用', '中：做成標準文件', '交付手冊與文件範本'],
        ['少發生 × 判斷多', '人決定，AI 只提供資料', '低：不產品化', '人工關卡與顧問判斷（我們的專業）']]
grid(s, 0.45, 1.45, [2.6, 3.0, 2.95, 3.9], ['象限', '怎麼做', '產品化潛力', '對應的產品元件'], ROWS, 0.82, size=11.5, hot=(0,))
callout(s, 5.0, '意涵：', '工具包從左上角開始做，最快看到效果、最容易重複賣；右下角是客戶最在意的控制點，也是顧問價值所在，兩者一起賣才完整。', h=0.85)

# ───────── 章節 03 ─────────
section('03', '市場驗證', '同業已經在賣什麼？我們從哪裡切入才有勝算？', ['同業的四種產品化型態', '我們的三個切入點與證據'])

# ───────── 8 市場 ─────────
s = content('市場驗證：同業已經把保險顧問服務做成產品，主流是「預建資產＋顧問導入」')
section_tag(s, '市場驗證', 2)
ROWS = [['加速器＋顧問導入', '國際會計師事務所的保險業代理型 AI 藍圖與預建代理；IFRS 17 加速器（業者宣稱 90 天內產出）', '最接近我們的模式；客戶會拿來比較', '採用：預建資產＋固定範圍與期程'],
        ['核心系統＋AI 工具組', '系統整合商的核心系統 AI 工具組（業者宣稱 300 個以上保險加速器）', '和核心系統轉型主題直接重疊', '聚焦台灣在地核心系統'],
        ['軟體／數據產品', '國際保險顧問的定價與核保軟體；策略顧問的保險市場資料庫訂閱', '產品化程度最高，但維護成本高', '長期方向，先不做軟體'],
        ['方法論與藍圖', '策略顧問的 AI-First 藍圖、代理型 AI 核心現代化方法', '還不是可單獨購買的產品', '我們直接落地到試點']]
grid(s, 0.45, 1.45, [2.3, 4.6, 2.85, 2.7], ['產品化型態', '代表做法', '對我們的意思', '我們怎麼做'], ROWS, 0.9, size=10.5, hot=(0,))
note(s, '具體公司、產品與來源見附錄 A。資料為業者公開資訊，未經獨立驗證；「宣稱」數字為業者自述。', y=5.6)

# ───────── 9 切入點 ─────────
s = content('我們的切入點：在地法規、在地核心系統、可稽核的關卡綁在一起，國際方案最難直接套用')
section_tag(s, '市場驗證', 2)
DIFF = [('台灣保險法規內建', '委外辦法、公會 AI 與雲端自律規範、金管會 AI 指引，整理成開發流程的檢查清單。', '證據：法規檢查清單草稿已完成（附錄 B）'),
        ('懂保險業務與在地系統', '理賠、核保、核心系統、IFRS 17 的業務與系統經驗，試點直接用客戶真正在痛的流程。', '證據：現行流程與 IT 實作成果（待補）'),
        ('可稽核的人工關卡', '五個關卡誰核准、何時核准都查得到；寫程式的 AI 不能核准自己的程式。', '證據：關卡設計（第 7 頁）')]
for i, (h, t, ev) in enumerate(DIFF):
    x = 0.45 + i * 4.18
    circle(s, x, 1.5, 0.55, str(i + 1), size=14)
    tb(s, x + 0.65, 1.5, 3.4, 0.55, h, size=14, bold=True, color=OR, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x, 2.15, 4.05, 1.6, t, size=12, fill=LG, margin=0.16)
    tb(s, x, 3.8, 4.05, 0.7, ev, size=11, bold=True, color=OR, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
callout(s, 4.8, '為什麼國際方案難直接套用（推論）：', '預建代理與藍圖以國際市場為主，套到台灣客製化的舊核心、繁中條款與在地法規時，仍需大量調整。', h=0.75)
note(s, '「難直接套用」為依公開資料的推論，僅供內部討論。', y=5.65)

# ───────── 章節 04 ─────────
section('04', '建議', '我們要賣什麼、怎麼收費、怎麼證明價值？', ['入門診斷加三種合作方式與對外策略', '試點成功標準'])

# ───────── 10 方案 ─────────
s = content('建議：一個入門診斷加三種合作方式，對外先推「診斷 → 陪你做試點」')
section_tag(s, '建議', 3)
ROWS = [['準備度診斷', '瓶頸地圖、法規檢查、試點建議、導入藍圖', '固定價', '對外主推', '低門檻，客戶容易說好'],
        ['陪你做｜試點陪跑＋教學', '完成的試點、成效報告、範本與審查規則、教材', '每個試點固定價，可加里程碑', '對外主推', '看得到成效，最能證明價值'],
        ['你自己做｜工具包', '範本、規範規則、關卡設定、交付手冊', '授權或年費', '客戶問了再談', '要先從試點累積成熟的資產'],
        ['幫你做｜駐點開發', '程式、文件與完整紀錄', '人月或按交付項目', '客戶問了再談', '偏人力派遣；審計客戶有獨立性限制']]
grid(s, 0.45, 1.45, [2.6, 3.7, 2.3, 1.55, 2.3], ['方案', '交付', '收費模式（建議）', '對外策略', '原因'], ROWS, 0.82, size=11, hot=(0, 1))
callout(s, 5.3, '計價原則：', '按「每次診斷」「每個試點」等任務或成果計價，不按工時，讓收入不再只靠人力增加。金額與期程今天決定。', h=0.7)
note(s, f'三種合作方式的概念出自 {SMR}（Do it yourself／with me／for me）。', y=6.1)

# ───────── 11 試點成功標準 ─────────
s = criteria_page('價值怎麼證明：試點前雙方講好三個指標和合格線，達標才擴大')
section_tag(s, '建議', 3)
note(s, '對外不先承諾效益數字；指標可依客戶痛點替換，以三個為限。', y=5.85)

# ───────── 章節 05 ─────────
section('05', '下一步', '各團隊要投入什麼？今天要拍板什麼？', ['對各團隊的意義', '要拍板的五件事與接下來三步'])

# ───────── 12 對各方的意義 ─────────
s = content('對每個人的意義：這個方案讓每個團隊都得到想要的東西，投入也講清楚')
section_tag(s, '下一步', 4)
ROWS = [['管理層', '一項新的可銷售服務；按任務或成果計價，不只靠人力成長', '核定價格區間與人力'],
        ['IT 團隊', '現行流程成為產品核心；實作成果成為對外實績', '提供實作成果或現場示範；確認四象限分類'],
        ['業務與 BA', '主導需求分析這一步；需求整理有固定方法與範本', '起草範本與法規檢查清單'],
        ['風險與品質', '關卡、紀錄與法規檢查清單讓 AI 使用可控、可稽核', '審閱法規檢查清單；確認獨立性原則'],
        ['對外溝通（短影片）', '影片直接用這套流程與示範畫面，內容有實證', '提供示範錄影素材']]
grid(s, 0.45, 1.45, [2.4, 5.6, 4.45], ['團隊', '得到什麼', '需要投入什麼'], ROWS, 0.7, size=12)
callout(s, 5.55, '共同點：', '同一個試點，同時產出工具包第一版、對外實績、影片示範與成效數據，一次投入、多處使用。', h=0.7)

# ───────── 13 決策 ─────────
s = content('今天要拍板的五件事，以及接下來三步')
section_tag(s, '下一步', 4)
DEC = [['四象限分類是否同意', 'IT 團隊', '今天'],
       ['試點案例：內部真實專案，或理賠進度查詢示範', 'IT 團隊、BA', '今天'],
       ['對外先推「診斷 → 陪你做試點」', '管理層', '今天'],
       ['價格區間、負責人與人力', '管理層', '今天'],
       ['法規檢查清單交付審閱；是否承諾「未達標不擴大」', '風險與品質、管理層', '今天']]
grid(s, 0.45, 1.45, [6.0, 2.6, 1.3], ['要決定的事', '負責', '期限'], DEC, 0.62, size=11.5, first_bold=False)
NX = [('1　補證據', '實作成果或示範；法規清單審閱'), ('2　定案', '客戶版簡報與價格定稿'), ('3　第一場工作坊', '對齊目標、挑候選試點')]
for i, (h, t) in enumerate(NX):
    y = 1.45 + i * 1.12
    tb(s, 10.55, y, 2.35, 1.0, [[(h, {'bold': True, 'size': 12.5, 'color': WHITE})], [(t, {'color': WHITE, 'size': 10.5})]], fill=[OR, OR2, INK][i], anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
callout(s, 5.15, '原則：', '客戶資料不放進外部 AI 工具；AI 協助產出的內容對外前一定人工確認；審計客戶推案前先確認獨立性。', h=0.7)

# ───────── 附錄轉場 ─────────
section('A', '附錄', '正文的資料從哪裡來？', ['同業具體產品與來源連結', '保險業 AI 開發法規檢查清單（草稿）'])

# ───────── 14 附錄 A ─────────
s = content('附錄 A｜同業具體產品與來源（業者公開資訊）')
ROWS = [('Deloitte', 'Agentic AI Blueprint for Insurance', '加速器＋導入', 'https://www.deloitte.com/global/en/what-we-do/capabilities/agentic-ai-blueprint-for-insurance.html'),
        ('Deloitte', 'IFRS 17 Regulatory Assist（與 SAS 合作）', '加速器', 'https://sas.com/nl_nl/news/press-releases/2022/may/deloitte-partners-with-sas-to-power-ifrs-17-accelerator-solution.html'),
        ('KPMG', 'IFRS 17 Express, powered by SAS', '加速器', 'https://sas.com/en_ca/news/press-releases/2022/september/pvi-insurance-deploys-kpmg-ifrs-17-express.html'),
        ('Capgemini', 'Guidewire 生成式 AI／代理工具組', '核心系統＋AI', 'https://www.capgemini.com/wp-content/uploads/2026/10/GW-datasheet_Final-version.pdf'),
        ('WTW', 'Radar 5（保險定價與核保軟體）', '軟體產品', 'https://www.wtwco.com/en-gh/news/2025/10/wtw-launch-of-radar-5-with-gen-ai-capability-marks-a-major-milestone-in-insurance-pricing'),
        ('McKinsey', 'Global Insurance Pools（訂閱資料庫）', '數據產品', 'https://www.mckinsey.com/industries/financial-services/how-we-help-clients/insurance/global-insurance-pools'),
        ('McKinsey', '代理型 AI 核心現代化', '方法論', 'https://www.mckinsey.com/industries/financial-services/our-insights/can-agentic-ai-finally-modernize-core-technologies-in-insurance'),
        ('BCG', 'AI-First 產險公司藍圖', '方法論', 'https://www.bcg.com/assets/2026/executive-perspectives-ai-first-companies-win-the-future-property-casualty-insurance.pdf')]
xx = [0.45, 2.0, 6.0, 7.75]
for j, h in enumerate(['公司', '產品／方案', '型態', '來源（可點擊）']):
    tb(s, xx[j], 1.45, (xx[j + 1] if j < 3 else 12.9) - xx[j] - 0.04, 0.36, h, size=10.5, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
for i, (c, prod, kind, url) in enumerate(ROWS):
    y = 1.85 + i * 0.5; bg = LG if i % 2 == 0 else WHITE
    tb(s, xx[0], y, 1.51, 0.46, c, size=10.5, bold=True, color=OR, fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
    tb(s, xx[1], y, 3.96, 0.46, prod, size=10.5, fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
    tb(s, xx[2], y, 1.71, 0.46, kind, size=10.5, fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
    tb(s, xx[3], y, 5.11, 0.46, [[(url.split('/')[2] + ' 原文', {'link': url})]], size=10.5, fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
note(s, '完整盤點見 research/productization/保險顧問產品化服務盤點.md。', y=6.0)

# ───────── 15 附錄 B ─────────
checklist_page('附錄 B｜保險業 AI 開發法規檢查清單（草稿）',
               '草稿，待風險與品質及法遵審閱；條文以主管機關公告為準，非法律意見。')

p.save(OUT)
print(OUT, len(p.slides), 'slides')
