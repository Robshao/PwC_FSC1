# AI Agent 協作開發流程：內部會議精簡版（9 頁）
# 以 Rex v3 五階段架構為基礎，說明為什麼要產品化，並沿用 AI 原生 SDLC 簡報的「直線變循環」圖；完整版見 build_framework_update.py
# 用法：python3 build_framework_brief.py <範本.pptx> <輸出.pptx> <手冊.pdf>
#   手冊的圖在建置時從 PDF 裁切，不放進版本控制
import sys, copy, subprocess, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageChops

TPL, OUT, PDF = sys.argv[1], sys.argv[2], sys.argv[3]
OR = RGBColor(0xFD, 0x51, 0x08); OR2 = RGBColor(0xFE, 0x7C, 0x39); TINT = RGBColor(0xFF, 0xE8, 0xD4)
# 範本佈景主題色：tx1 黑、lt2 淺灰、accent6 線條灰、accent4 中灰
INK = RGBColor(0x00, 0x00, 0x00); GREY = RGBColor(0x59, 0x59, 0x59); LG = RGBColor(0xEB, 0xEB, 0xEB); LINE = RGBColor(0xCB, 0xD1, 0xD6)
MID = RGBColor(0xA1, 0xA8, 0xB3); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# 範本字型：標題 Georgia、內文 Arial、中文一律微軟正黑體
FONT = 'Arial'; HEAD_FONT = 'Georgia'; EA_FONT = '微軟正黑體'
SYMBOLS = set('●◐○✓✗→↑①②③④⑤⑥⑦⚠')
FOOT = '內部會議精簡版｜以 Rex《AI Agent 協作開發流程與產出範本》v3 為基礎｜AI 協助產出，需人工確認後再對外'
V3 = 'Rex《AI Agent 協作開發流程與產出範本》v3'
PB = 'Anthropic《The AI-Native SDLC Playbook》'
SMR = 'Sawhney, MIT Sloan Management Review, 2024'

def style_font(r, head=False):
    # Arial 沒有 ●◐○ 等符號，這類文字改用中文字型顯示
    r.font.name = EA_FONT if SYMBOLS & set(r.text) else (HEAD_FONT if head else FONT)
    rPr = r._r.get_or_add_rPr()
    for e in rPr.findall(qn('a:ea')): rPr.remove(e)
    rPr.find(qn('a:latin')).addnext(rPr.makeelement(qn('a:ea'), {'typeface': EA_FONT}))

p = Presentation(TPL)
lst = p.slides._sldIdLst
for sid in list(lst):
    p.part.drop_rel(sid.rId); lst.remove(sid)
L = {l.name: l for l in p.slide_masters[2].slide_layouts}
# 移除範本遺留的投影片章節（PowerPoint 左側的章節名稱）
for ext in list(p.part._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}ext')):
    if ext.find('{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst') is not None:
        ext.getparent().remove(ext)

def set_ph(s, idx, text, size=None, bold=None):
    for ph in s.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            for para in ph.text_frame.paragraphs:
                for r in para.runs:
                    style_font(r, head=idx == 0)
                    if size: r.font.size = Pt(size)
                    if bold is not None: r.font.bold = bold
            return ph

def tb(s, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, margin=0.06, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'): setattr(tf, f'margin_{side}', Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg in segs:
            t, st = (seg, {}) if isinstance(seg, str) else seg
            r = para.add_run(); r.text = t
            style_font(r); f = r.font; f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color)
        para.space_after = Pt(2)
    return sh

def circle(s, x, y, d, text, fill=OR, size=12):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    r = para.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; style_font(r)

def hline(s, x1, y1, x2, y2, color=LINE, width=1.0):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width)

def arrow(s, x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, fill=OR):
    a = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = fill; a.line.fill.background(); a.shadow.inherit = False

def footer(s):
    # python-pptx 新增投影片不會帶入頁尾與頁碼，從版面複製（頁碼為自動欄位）
    have = {ph.placeholder_format.idx for ph in s.placeholders}
    for ph in s.slide_layout.placeholders:
        if ph.placeholder_format.idx in (3, 4) and ph.placeholder_format.idx not in have:
            s.shapes._spTree.append(copy.deepcopy(ph._element))
    set_ph(s, 3, FOOT)

def content(title):
    s = p.slides.add_slide(L['只有標題'])
    set_ph(s, 0, title, 20, bold=True); footer(s)
    return s

def note(s, text, y=6.66):
    tb(s, 0.45, y, 12.45, 0.32, text, size=9, color=GREY, margin=0.03)

def section(num, title, ask, see):
    # 章節轉場頁：沿用範本「章節標題」版面，左上方說明這一章要回答什麼、會看到什麼
    s = p.slides.add_slide(L['章節標題'])
    set_ph(s, 0, title, 30); set_ph(s, 1, num, 180); footer(s)
    tb(s, 0.45, 0.7, 7.9, 2.75, [[('這一章要回答', {'bold': True, 'color': OR, 'size': 13})], [(ask, {'size': 17, 'bold': True})],
                                 [('', {'size': 6})], [('接下來你會看到', {'bold': True, 'color': OR, 'size': 13})]] + [[('・' + x, {'size': 14})] for x in see],
       fill=WHITE, margin=0.22)
    return s

def grid(s, x, y, cw, head, rows, rh, size=10.5, first_bold=True, head_h=0.4, hot=()):
    # 以文字框排成的表格：表頭橘底白字、隔列淺灰；hot 為要強調的列
    xx = x
    for j, h in enumerate(head):
        tb(s, xx, y, cw[j] - 0.04, head_h, h, size=size, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
        xx += cw[j]
    for i, row in enumerate(rows):
        yy = y + head_h + 0.04 + i * rh; xx = x
        bg = TINT if i in hot else (LG if i % 2 == 0 else WHITE)
        for j, v in enumerate(row):
            first = j == 0 and first_bold
            tb(s, xx, yy, cw[j] - 0.04, rh - 0.04, v, size=size, bold=first, color=OR if first else INK,
               fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
            xx += cw[j]

def callout(s, y, label, text, h=0.62, size=11.5):
    tb(s, 0.45, y, 0.1, h, '', fill=OR)
    tb(s, 0.55, y, 12.35, h, [[(label, {'bold': True, 'color': OR}), (text, {})]], size=size, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.16)


# ───────── 從手冊 PDF 裁切「直線變循環」圖（與 build_sdlc_deck.py 相同裁切範圍） ─────────
TMP = tempfile.mkdtemp()
subprocess.run(['pdftoppm', '-r', '220', '-f', '5', '-l', '5', '-png', '-singlefile', PDF, f'{TMP}/p5'], check=True)
im = Image.open(f'{TMP}/p5.png').convert('RGB'); w, h = im.size
im = im.crop((int(0.09 * w), int(0.335 * h), int(0.91 * w), int(0.775 * h)))
box = ImageChops.difference(im, Image.new('RGB', im.size, (255, 255, 255))).convert('L').point(lambda v: 255 if v > 12 else 0).getbbox()
FIG2 = f'{TMP}/fig2.png'
im.crop((max(box[0] - 6, 0), max(box[1] - 6, 0), min(box[2] + 6, im.size[0]), min(box[3] + 6, im.size[1]))).save(FIG2)

def pic_fit(s, path, x, y, mw, mh):
    # 依原圖比例放入，不超過 mw × mh
    iw, ih = Image.open(path).size
    pw = min(mw, mh * iw / ih); ph = pw * ih / iw
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(pw), Inches(ph))
    return pw, ph

# ───────── 1 封面 ─────────
s = p.slides.add_slide(L['標題投影片'])
set_ph(s, 0, 'AI Agent 協作開發流程：\n內部會議精簡版', 28)
set_ph(s, 1, 'FSC1 內部會議｜為什麼要做成產品｜2026 年 10 月', 12)

# ───────── 2 一頁看懂 ─────────
s = content('一頁看懂：把 Rex 的開發流程做成可以重複提供的服務，先在自己的案例跑通，再帶客戶做')
KEY = [('做什麼', '以 Rex v3 的五個階段為骨架，加上指定核准人的關卡、每階段的交接檔案，以及上線後的回饋，把直線流程變成循環。'),
       ('為什麼做成產品', '只接專案，營收要靠多請人才能成長；同業也已經把保險顧問服務做成產品。把流程整理成工具包，每做一個案子都留下可以重複使用的東西。'),
       ('怎麼開始', '用一個案例走完五個階段，同時產出工具包第一版、五頁案例頁和第 3 集影片的示範畫面。對外主推「陪你做一個試點」。')]
for i, (h, t) in enumerate(KEY):
    y = 1.45 + i * 1.45
    circle(s, 0.45, y + 0.3, 0.6, str(i + 1), size=16)
    tb(s, 1.2, y, 1.9, 1.25, h, size=14, bold=True, color=OR, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 3.1, y, 4.9, 1.25, t, size=12, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
DEC = ['五頁案例頁用哪一個專案（IT 團隊的真實專案，或理賠進度自助查詢示範）',
       '需求分析由 BA、IT、AI 怎麼分工',
       '工具包第一版：三份範本＋一條規範技能＋上線閘門',
       '誰錄五個階段的畫面、什麼時候錄']
tb(s, 8.3, 1.45, 4.6, 0.5, '今天要決定', size=14, bold=True, color=WHITE, fill=INK, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
tb(s, 8.3, 2.0, 4.6, 3.8, [[(f'{i + 1}　', {'bold': True, 'color': OR}), (t, {})] for i, t in enumerate(DEC)], size=12.5, fill=TINT, margin=0.18)
callout(s, 6.0, '一句話：', 'AI 讓每一步更快、每一步都有紀錄；要不要做、能不能上線由指定的人決定。我們把這套做法做成工具包和服務。', h=0.55)

# ───────── 3 為什麼要產品化 ─────────
s = content('為什麼要做成產品：只接專案會卡在人力，同業已經在做，而這套流程本來就適合產品化')
WHY = [('只接專案，成長卡在人力', '服務靠人力，毛利偏低；多賺一塊錢就要多請一個人。文章舉例：SaaS 新創估值約為年營收 6–8 倍，專案型服務只有 1–2 倍。'),
       ('同業已經在做', 'Deloitte 有保險業 AI 導入藍圖與預建代理，Capgemini 有核心系統 AI 工具組，Deloitte、KPMG 把 IFRS 17 做成加速器。客戶會拿這些來比較。'),
       ('這套流程本來就適合', '開發工作多半是「常發生、規則清楚」的活動，例如欄位清單、測試案例、文件。文章指出這類活動最適合做成產品。')]
tb(s, 0.45, 1.4, 6.1, 0.42, '三個理由', size=13, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
for i, (h, t) in enumerate(WHY):
    tb(s, 0.45, 1.88 + i * 1.2, 6.1, 1.14, [[(f'{i + 1}　{h}', {'bold': True, 'color': OR, 'size': 12.5})], t], size=11, fill=LG, margin=0.12)
tb(s, 6.8, 1.4, 6.1, 0.42, '我們有沒有條件：產品化需要的三種能力', size=13, bold=True, color=WHITE, fill=INK, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
CAP = [['領域專業', '保險流程、法規、核心系統、IFRS 17', '寫成技能與檢查清單'],
       ['自有數據', '研究、簡報、影片腳本', '範本、規範規則庫、專案紀錄，要從案例累積'],
       ['技術平台', 'IT 團隊已在用 AI 協作開發', '工具包第一版（範本、技能、閘門）']]
grid(s, 6.8, 1.88, [1.35, 2.35, 2.4], ['能力', '已經有', '還要補'], CAP, 1.08, size=11, head_h=0.36)
callout(s, 5.55, '提醒：', '產品化常要花好幾年，而且是從一個客製專案長出來的（文章案例從 2013 年的專案到 2015 年才成為產品）。同業也有把軟體產品賣掉的例子，維護成本不低。所以先做「陪你做」的試點和工具包，不直接做軟體。', h=0.88, size=11)
note(s, f'來源：{SMR}；同業資料見 research/productization/保險顧問產品化服務盤點.md（業者自述，未經獨立驗證）。估值倍數為文章所述。', y=6.5)

# ───────── 4 直線變循環 ─────────
s = content('從直線變循環：Rex 的五個階段放進循環，每一步交出檔案、由指定的人核准')
pw, ph = pic_fit(s, FIG2, 0.45, 1.4, 6.6, 3.9)
tb(s, 0.45, 1.45 + ph, pw, 0.3, '圖：Anthropic 手冊圖 2（同 AI 原生 SDLC 簡報第 6 頁）', size=9, color=GREY, margin=0.02)
x0 = 0.45 + pw + 0.3
ROWS = [['需求分析', '需求說明、Prototype', '① 產品負責人'],
        ['SA&D', '規格書、ERD、API', '② 產品負責人＋技術主管'],
        ['開發', '實作計畫、程式', '③ 工程師'],
        ['測試', '測試結果、報告', 'AI 先自行驗證'],
        ['上版／交付', '審查與核准紀錄', '④ 程式負責人 ⑤ 發布經理'],
        ['（新增）維運', '事件紀錄 → 新需求', '服務負責人']]
grid(s, x0, 1.4, [1.4, (12.9 - x0 - 1.4) * 0.5, (12.9 - x0 - 1.4) * 0.5], ['v3 階段', '交出的檔案', '誰核准'], ROWS, 0.56, size=10.5, head_h=0.36, hot=(5,))
callout(s, 5.55, '和 v3 的差別：', '核心原則「AI 先產出 → IT Review → 修正 → 定稿 → 才進下一階段」不變。新增兩點：每個關卡有指定核准人並留下紀錄；上線後的異常由 AI 先診斷，寫成新的需求說明，回到第一步。', h=0.85, size=11)
note(s, '左圖：左邊傳統的直線，六個階段依序交接；右邊 AI 原生的循環，AI 參與每個階段，人在上方發起、指揮與把關。', y=6.5)

# ───────── 5 矩陣 ─────────
s = content('哪些先交給 AI，也就是哪些先做成工具包：常發生、規則清楚的工作優先')
MX, MY, QW, QH = 1.15, 1.5, 3.6, 2.15
Q = [(0, 0, '常發生 × 判斷少', 'AI 先做，IT 確認', '欄位清單、Test Case、Test Data、Unit Test、註解與文件、Release Note', TINT),
     (1, 0, '常發生 × 判斷多', 'AI 起草，IT 逐項審查', '程式碼、API 規格、ERD、Code Review 風險檢查、測試失敗原因分析', LG),
     (0, 1, '少發生 × 判斷少', '做成範本，重複使用', 'Deploy Guide、User Guide、Migration Script、部署腳本', LG),
     (1, 1, '少發生 × 判斷多', '人決定，AI 只提供資料', '接受需求、架構選型、核准合併、授權上線、事件處置', WHITE)]
for c, r, h, rule, items, bg in Q:
    x, y = MX + c * (QW + 0.06), MY + r * (QH + 0.06)
    tb(s, x, y, QW, QH, '', fill=bg, line=LINE)
    tb(s, x + 0.12, y + 0.08, QW - 0.24, 0.62, [[(h, {'bold': True, 'size': 11, 'color': GREY})], [(rule, {'bold': True, 'size': 13, 'color': OR})]], margin=0.02)
    tb(s, x + 0.12, y + 0.76, QW - 0.24, QH - 0.84, items, size=11.5, margin=0.02)
tb(s, 0.45, MY, 0.65, QH, '常發生', size=11, bold=True, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
tb(s, 0.45, MY + QH + 0.06, 0.65, QH, '少發生', size=11, bold=True, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
tb(s, MX, MY + 2 * QH + 0.1, 2 * QW + 0.06, 0.3, '需要判斷：少 ─────────────→ 多', size=10.5, bold=True, color=GREY, align=PP_ALIGN.CENTER, margin=0.0)
tb(s, 8.75, 1.5, 4.15, 4.75, [[('和產品化的關係', {'bold': True, 'color': OR, 'size': 13})],
                              [('左上：', {'bold': True}), ('工具包第一版從這裡做。最常發生、最快看到效果，也最容易重複賣。', {})],
                              [('右上：', {'bold': True}), ('「陪你做」時由我們帶客戶審查。AI 讓審查變快，不會讓審查消失。', {})],
                              [('左下：', {'bold': True}), ('做成範本，放進工具包。', {})],
                              [('右下：', {'bold': True}), ('永遠留給客戶的人，對應五個人工關卡。', {})],
                              [('', {'size': 6})],
                              [('出處：', {'bold': True, 'color': OR}), (f'{SMR} 產品化第一步：把服務拆成活動，用「頻率」和「認知複雜度」評估。', {})]],
   size=11.5, fill=LG, margin=0.16)
note(s, f'任務清單取自 {V3} 各階段「AI 可以做什麼」；分類為提案，請 IT 團隊依實際專案調整。', y=6.5)

# ───────── 6 工具包第一版 ─────────
s = content('工具包第一版做什麼：把規則從 Prompt 寫進範本、技能與自動閘門，不靠個人功力')
LV = [('1', 'Prompt 六要素', 'v3 附錄已有。效果因人而異', MID),
      ('2', '範本與團隊說明檔', '需求說明、規格、計畫用固定格式；團隊慣例 AI 每次都先讀', OR),
      ('3', '技能（Skill）', '資安、個資、委外規範寫成檔案，自動套用；讓違規變少', OR),
      ('4', '自動閘門（Hook）', '沒有核准不能合併、不能上線，個人關不掉；讓違規幾乎做不到', OR)]
for i, (n, h, t, c) in enumerate(LV):
    x = 0.45 + i * 3.14
    tb(s, x, 1.4, 3.0, 0.5, f'{n}　{h}', size=12.5, bold=True, color=WHITE, fill=c, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
    tb(s, x, 1.95, 3.0, 1.15, t, size=11.5, fill=LG, margin=0.12)
    if i < 3: arrow(s, x + 3.0, 2.4, 0.14, 0.26)
V1 = ['需求說明、規格、實作計畫三份範本（第 2 層）',
      '一條保險業規範技能，例如「入口網站不能新增個資」（第 3 層）',
      '上線閘門：沒有核准不能合併、不能上線（第 4 層）',
      'Review Checklist 九項',
      '五階段案例：每階段三張截圖、一個數字']
tb(s, 0.45, 3.3, 6.1, 0.42, '第一版交付內容', size=12.5, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
tb(s, 0.45, 3.76, 6.1, 2.35, [[(f'{i + 1}　', {'bold': True, 'color': OR}), (t, {})] for i, t in enumerate(V1)], size=12, fill=TINT, margin=0.16)
CHK = [('v3 已有五項', '正確、完整、一致、安全、可追溯'),
       ('法遵', '有沒有用到個資？資料是否在可給 AI 讀的範圍內？'),
       ('職責分離', '寫程式的 AI 不能自己核准；關卡由指定的人簽'),
       ('測試證據', '有自動測試結果，案例綁需求編號'),
       ('紀錄完整', '誰核准、何時、套用哪版規範，查得到')]
tb(s, 6.8, 3.3, 6.1, 0.42, 'Review Checklist：v3 五項＋新增四項', size=12.5, bold=True, color=WHITE, fill=INK, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
for i, (h, t) in enumerate(CHK):
    tb(s, 6.8, 3.76 + i * 0.47, 1.45, 0.43, h, size=11, bold=True, color=OR if i else GREY, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
    tb(s, 8.3, 3.76 + i * 0.47, 4.6, 0.43, t, size=11, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
note(s, f'新增四項對應保險業 AI 自律規範與金管會 AI 指引的治理與問責原則。來源：{V3} 第 16、18 頁；{PB} 第 25–27、40–46 頁。', y=6.3)

# ───────── 7 怎麼賣 ─────────
s = content('怎麼賣：同一套流程包成三種方案，主推「陪你做一個試點」')
TIER = [('你自己做', '工具包', '範本、規範技能、自動閘門與 Review Checklist；適合已有 IT 團隊、只缺方法的客戶', MID),
        ('陪你做（主推）', '試點陪跑＋教學', '現況診斷、挑一個試點、帶團隊用真實需求走完五階段、試點前後比較成效', OR),
        ('幫你做', '駐點開發', 'FSC1 人力駐點用這套流程開發，交付程式、文件與完整紀錄，期末交接工具包', OR2)]
for i, (h, sub, t, c) in enumerate(TIER):
    x = 0.45 + i * 4.18
    tb(s, x, 1.4, 4.05, 0.55, [[(h, {'bold': True, 'size': 14}), ('　' + sub, {'size': 11})]], color=WHITE, fill=c, anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
    tb(s, x, 2.0, 4.05, 1.05, t, size=11, fill=LG if i != 1 else TINT, margin=0.12)
ROWS = [['1 評估潛力', '能做成產品到什麼程度？', '●', '已用矩陣分好（第 5 頁）'],
        ['2 決定程度', '客戶想要多少人力參與？', '◐', '三種方案草案，範圍待 Sam、Luke 決定'],
        ['3 服務組合', '給哪種客戶什麼方案？', '◐', '主推陪你做；差異：台灣法規、在地核心系統、可稽核的人工關卡'],
        ['4 建立能力', '領域專業、自有數據、技術平台', '○', '靠第一個案例累積範本、規則庫與紀錄'],
        ['5 帶領轉型', '計價、KPI、組織怎麼改？', '○', '報價考慮改成「每次診斷」「每個試點」固定價，不按人天']]
tb(s, 0.45, 3.25, 12.45, 0.36, '產品化五步驟：我們走到哪裡', size=12, bold=True, color=GREY, margin=0.02)
grid(s, 0.45, 3.62, [1.6, 3.4, 0.9, 6.55], ['步驟', '要回答的問題', '進度', '我們的狀況'], ROWS, 0.48, size=10.5, head_h=0.34)
note(s, f'●完成　◐部分　○尚未開始。三種方案與五步驟出自 {SMR}。對外不點名同業、影片不提價格。', y=6.5)

# ───────── 8 下一步 ─────────
s = content('下一步：一個案例同時產出工具包、案例頁和影片示範')
ROWS = [['選定案例並錄製五階段畫面（示範資料一律用假資料）', 'Rex、Kenny', '會議決定'],
        ['填寫 IT 團隊實績填寫表', 'Kenny／IT 團隊', '會議決定'],
        ['三份範本、規範技能、Review Checklist 初版', 'BA 起草、Rex 確認', '會議決定'],
        ['工具包名稱與範圍、三種方案報價', 'Sam、Luke', '會議決定'],
        ['第 3 集影片腳本檢視（第③段示範用同一批錄影）', '影片組、Sam', '10/28'],
        ['獨立性檢查', '影片組、R&Q', '11/04']]
grid(s, 0.45, 1.4, [7.6, 2.85, 2.0], ['事項', '負責', '期限'], ROWS, 0.55, size=11.5, first_bold=False, hot=(0,))
tb(s, 0.45, 5.25, 6.1, 1.0, [[('一次錄、三處用', {'bold': True, 'color': OR, 'size': 12})], '五頁案例頁、第 3 集影片第③段示範、對客簡報「為什麼是 PwC FSC1」的實績欄。'], size=11.5, fill=TINT, margin=0.14)
tb(s, 6.8, 5.25, 6.1, 1.0, [[('所內規定', {'bold': True, 'color': OR, 'size': 12})], '客戶資料不放外部 AI 工具；AI 產出對外前要人工確認；審計客戶先和 R&Q 確認獨立性。'], size=11.5, fill=LG, margin=0.14)
note(s, '負責人與期限為提案，以會議決定為準。完整說明見 24 頁更新版。', y=6.4)

# ───────── 9 來源 ─────────
s = content('附錄｜來源')
SRCS = [('上週版本', f'{V3}（19 頁）。五階段架構、IT／AI 分工、Review Checklist、Prompt 六要素。'),
        ('流程與圖', f'Louis Claxton，{PB}，2026，圖 2「The line and the loop」。白話說明見 research/ai-native-sdlc/AI原生SDLC_白話說明.pptx。'),
        ('產品化', 'Mohanbir Sawhney，〈How to Turn Professional Services Into Products〉，MIT Sloan Management Review，2024。整理見 research/productization/MIT-SMR_專業服務產品化_文章報告.md。'),
        ('同業盤點', 'research/productization/保險顧問產品化服務盤點.md（業者新聞稿與產品頁，含連結）。'),
        ('完整版', 'research/ai-native-sdlc/AI協作開發流程_小組會議更新版.pptx（24 頁，含法規與成效指標）。')]
for i, (h, t) in enumerate(SRCS):
    y = 1.4 + i * 0.95
    tb(s, 0.45, y, 2.0, 0.85, h, size=12, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
    tb(s, 2.5, y, 10.4, 0.85, t, size=11.5, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)

p.save(OUT)
print(OUT, len(p.slides), 'slides')
