# AI Agent 協作開發流程：小組會議更新版
# 以 Rex《AI Agent 協作開發流程與產出範本》v3 的五階段架構為基礎，補上 Anthropic 手冊、MIT SMR 產品化框架、同業盤點與法規
# 用法：python3 build_framework_update.py <範本.pptx> <輸出.pptx>
import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TPL, OUT = sys.argv[1], sys.argv[2]
OR = RGBColor(0xFD, 0x51, 0x08); OR2 = RGBColor(0xFE, 0x7C, 0x39); TINT = RGBColor(0xFF, 0xE8, 0xD4)
# 範本佈景主題色：tx1 黑、lt2 淺灰、accent6 線條灰、accent4 中灰
INK = RGBColor(0x00, 0x00, 0x00); GREY = RGBColor(0x59, 0x59, 0x59); LG = RGBColor(0xEB, 0xEB, 0xEB); LINE = RGBColor(0xCB, 0xD1, 0xD6)
MID = RGBColor(0xA1, 0xA8, 0xB3); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# 範本字型：標題 Georgia、內文 Arial、中文一律微軟正黑體
FONT = 'Arial'; HEAD_FONT = 'Georgia'; EA_FONT = '微軟正黑體'
SYMBOLS = set('●◐○✓✗→↑①②③④⑤⑥⑦⚠')
FOOT = '內部討論｜以 Rex《AI Agent 協作開發流程與產出範本》v3 為基礎更新｜AI 協助產出，需人工確認後再對外'
V3 = 'Rex《AI Agent 協作開發流程與產出範本》v3'
PB = 'Anthropic《The AI-Native SDLC Playbook》'
SMR = 'Sawhney, MIT Sloan Management Review, 2024'

def style_font(r, head=False):
    # Arial 沒有 ●◐○ 等符號，這類文字改用中文字型顯示
    r.font.name = EA_FONT if SYMBOLS & set(r.text) else (HEAD_FONT if head else FONT)
    rPr = r._r.get_or_add_rPr()
    for e in rPr.findall(qn('a:ea')): rPr.remove(e)
    rPr.find(qn('a:latin')).addnext(rPr.makeelement(qn('a:ea'), {'typeface': EA_FONT}))

p = Presentation(TPL)
lst = p.slides._sldIdLst
for sid in list(lst):
    p.part.drop_rel(sid.rId); lst.remove(sid)
L = {l.name: l for l in p.slide_masters[2].slide_layouts}
# 移除範本遺留的投影片章節（PowerPoint 左側的章節名稱）
for ext in list(p.part._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}ext')):
    if ext.find('{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst') is not None:
        ext.getparent().remove(ext)

def set_ph(s, idx, text, size=None, bold=None):
    for ph in s.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            for para in ph.text_frame.paragraphs:
                for r in para.runs:
                    style_font(r, head=idx == 0)
                    if size: r.font.size = Pt(size)
                    if bold is not None: r.font.bold = bold
            return ph

def tb(s, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, margin=0.06, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'): setattr(tf, f'margin_{side}', Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg in segs:
            t, st = (seg, {}) if isinstance(seg, str) else seg
            r = para.add_run(); r.text = t
            style_font(r); f = r.font; f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color)
        para.space_after = Pt(2)
    return sh

def circle(s, x, y, d, text, fill=OR, size=12):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    r = para.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; style_font(r)

def hline(s, x1, y1, x2, y2, color=LINE, width=1.0):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width)

def arrow(s, x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, fill=OR):
    a = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = fill; a.line.fill.background(); a.shadow.inherit = False

def footer(s):
    # python-pptx 新增投影片不會帶入頁尾與頁碼，從版面複製（頁碼為自動欄位）
    have = {ph.placeholder_format.idx for ph in s.placeholders}
    for ph in s.slide_layout.placeholders:
        if ph.placeholder_format.idx in (3, 4) and ph.placeholder_format.idx not in have:
            s.shapes._spTree.append(copy.deepcopy(ph._element))
    set_ph(s, 3, FOOT)

def content(title):
    s = p.slides.add_slide(L['只有標題'])
    set_ph(s, 0, title, 20, bold=True); footer(s)
    return s

def note(s, text, y=6.66):
    tb(s, 0.45, y, 12.45, 0.32, text, size=9, color=GREY, margin=0.03)

def section(num, title, ask, see):
    # 章節轉場頁：沿用範本「章節標題」版面，左上方說明這一章要回答什麼、會看到什麼
    s = p.slides.add_slide(L['章節標題'])
    set_ph(s, 0, title, 30); set_ph(s, 1, num, 180); footer(s)
    tb(s, 0.45, 0.7, 7.9, 2.75, [[('這一章要回答', {'bold': True, 'color': OR, 'size': 13})], [(ask, {'size': 17, 'bold': True})],
                                 [('', {'size': 6})], [('接下來你會看到', {'bold': True, 'color': OR, 'size': 13})]] + [[('・' + x, {'size': 14})] for x in see],
       fill=WHITE, margin=0.22)
    return s

def grid(s, x, y, cw, head, rows, rh, size=10.5, first_bold=True, head_h=0.4, hot=()):
    # 以文字框排成的表格：表頭橘底白字、隔列淺灰；hot 為要強調的列
    xx = x
    for j, h in enumerate(head):
        tb(s, xx, y, cw[j] - 0.04, head_h, h, size=size, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
        xx += cw[j]
    for i, row in enumerate(rows):
        yy = y + head_h + 0.04 + i * rh; xx = x
        bg = TINT if i in hot else (LG if i % 2 == 0 else WHITE)
        for j, v in enumerate(row):
            first = j == 0 and first_bold
            tb(s, xx, yy, cw[j] - 0.04, rh - 0.04, v, size=size, bold=first, color=OR if first else INK,
               fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
            xx += cw[j]

def callout(s, y, label, text, h=0.62, size=11.5):
    tb(s, 0.45, y, 0.1, h, '', fill=OR)
    tb(s, 0.55, y, 12.35, h, [[(label, {'bold': True, 'color': OR}), (text, {})]], size=size, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.16)

# ───────── 1 封面 ─────────
s = p.slides.add_slide(L['標題投影片'])
set_ph(s, 0, 'AI Agent 協作開發流程：\n更新版', 28)
set_ph(s, 1, 'FSC1 小組會議｜以 Rex v3 為基礎的更新版｜2026 年 10 月', 12)

# ───────── 2 議程 ─────────
s = content('議程')
AG = [('', '一頁看懂', '這一版保留什麼、補了什麼、今天要決定什麼', '3'),
      ('01', '上週版本回顧', 'Rex v3 已經做到的，以及還空著的地方', '4–5'),
      ('02', '流程更新：從直線到循環', '五階段對照 Anthropic 手冊、人工關卡、需求分析分工、哪些工作交給 AI', '6–10'),
      ('03', 'AI 協作怎麼落地', '從寫好 Prompt 到把規則寫進工具、Review Checklist、案例頁怎麼填', '11–14'),
      ('04', '治理、法規與成效', '保險業法規與所內規定、每個階段要量的指標', '15–17'),
      ('05', '對外：服務化與影片', '三種合作方式、同業比較、第 3 集影片怎麼用這套流程', '18–21'),
      ('06', '下一步', '今天要決定的事與分工', '22–23'),
      ('附錄', '來源', '本簡報用到的資料', '24')]
for i, (n, h, d, pg) in enumerate(AG):
    y = 1.4 + i * 0.66
    if n: circle(s, 0.6, y + 0.06, 0.48, n, fill=OR if n != '附錄' else MID, size=10 if n == '附錄' else 12)
    tb(s, 1.3, y, 4.1, 0.56, h, size=15, bold=True, color=OR if n == '06' else INK, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 5.45, y, 6.2, 0.56, d, size=11.5, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 11.7, y, 1.2, 0.56, pg, size=15, bold=True, color=OR, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    hline(s, 1.3, y + 0.61, 12.9, y + 0.61)
note(s, '右側數字為頁碼。')

# ───────── 3 一頁看懂 ─────────
s = content('一頁看懂：Rex v3 的架構保留，這一版補上「誰核准、規則怎麼守、怎麼證明有效」')
COLS = [('保留', OR, ['五個階段：需求分析、SA&D、開發、測試、上版／交付',
                      '每階段的 Input、IT Action、AI 協作與 Review Gate',
                      '核心原則：AI 先產出 → IT Review → 修正 → 定稿 → 才進下一階段',
                      'Review Checklist 與 Prompt 六要素']),
        ('補強', OR2, ['每個 Gate 指定核准人，並留下紀錄',
                       '加上「維運」：上線後的問題回到需求，流程變循環',
                       '用「多常發生 × 需要判斷」決定哪些工作先交給 AI',
                       '規則從 Prompt 寫進範本、技能與自動閘門',
                       '法規、稽核與成效指標']),
        ('今天要決定', INK, ['用哪一個案例貫穿五個階段的案例頁',
                            '需求分析由 BA 和 IT 怎麼分工',
                            '工具包要先做哪幾樣',
                            '誰錄案例畫面、什麼時候錄'])]
for i, (h, c, items) in enumerate(COLS):
    x = 0.45 + i * 4.18
    tb(s, x, 1.42, 4.0, 0.5, h, size=14, bold=True, color=WHITE, fill=c, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
    tb(s, x, 1.97, 4.0, 3.55, [[('・' + t, {})] for t in items], size=13.5, fill=LG if i < 2 else TINT, margin=0.16)
callout(s, 5.75, '一句話：', 'AI 負責把每一步做得更快，每一步都留下紀錄；要不要做、能不能上線，由指定的人決定。我們把這套做法整理成工具包，先在自己的案例上跑一次，再帶客戶做。', h=0.8, size=12)

# ───────── 章節 01 ─────────
section('01', '上週版本\n回顧', 'Rex v3 已經做到哪裡？還有哪些地方是空的？', ['v3 的八個項目逐一盤點', '每一項在這一版怎麼處理'])

# ───────── 5 v3 盤點 ─────────
s = content('Rex v3 盤點：流程與分工已經完整，空著的是案例、需求分析分工、上線後回饋與治理')
ROWS = [['五階段流程與 Review Gate', '五個階段、每階段一個 Gate', '●', '保留；每個 Gate 補上「誰核准」與留下的紀錄'],
        ['IT 與 AI 分工', '每階段都有 Input、IT Action、AI Assist', '●', '保留；用「多常發生 × 需要判斷」排出先交給 AI 的工作'],
        ['Review Checklist', '正確、完整、一致、安全、可追溯五項', '◐', '加上法遵、職責分離、測試證據、紀錄完整四項'],
        ['Prompt 撰寫六要素', '放在附錄，靠每個人寫好 Prompt', '◐', '保留當第一層；往範本、技能、自動閘門做，不只靠個人功力'],
        ['實際協作案例（第 7、9、11、13、15 頁）', '五頁都是待貼截圖的空白框', '○', '用一個案例貫穿五個階段，同一批畫面也給影片示範使用'],
        ['需求分析的分工', '標註「待與 BA 討論」', '○', '提出 BA、IT、AI 三方分工建議（第 9 頁）'],
        ['上線之後', '上版／交付是終點', '○', '加上維運：問題寫成新的需求，回到第一步'],
        ['法規、稽核、成效', '尚未涵蓋', '○', '新增保險業法規、所內規定與每階段指標']]
grid(s, 0.45, 1.4, [3.3, 3.6, 0.9, 4.65], ['項目', 'v3 現況', '完成度', '這一版怎麼處理'], ROWS, 0.56, size=10.5, hot=(4, 5))
note(s, f'●完整　◐部分　○尚未處理。淺橘底為最需要 IT 團隊與 BA 一起補的兩項。依據：{V3}。', y=6.62)

# ───────── 章節 02 ─────────
section('02', '流程更新：\n從直線到循環', '五個階段要怎麼調整，才能讓每一步都有人負責、有紀錄可查？', ['v3 五階段對照 Anthropic 手冊的六個階段', '五個人工關卡與交接檔案', '需求分析 BA × IT × AI 分工', '哪些工作先交給 AI'])

# ───────── 7 對照 ─────────
s = content('五階段對照 Anthropic 手冊：每一階段都留下一份下一階段直接接手的檔案，關卡由指定的人核准')
ROWS = [['需求分析', '規劃', '需求說明（intent.md）\nPrototype、欄位清單', '① 接受需求：產品負責人', 'BA 主導和業務對焦；AI 追問範圍與限制'],
        ['SA&D', '設計', '規格書（spec.md）\nERD、API 規格、Test Plan', '② 確認規格：產品負責人（高風險加技術主管）', '規格自動套用資安、法遵規範，衝突處標給規範負責人'],
        ['開發', '開發', '實作計畫（plan.md）\n程式、Unit Test', '③ 核准計畫：工程師（高風險找架構師）', 'v3 已提「先 plan 再 code」；計畫核准前 AI 不能改程式'],
        ['測試', '測試', 'Test Case、測試結果\nBug List、Test Report', 'AI 先自行驗證，全部通過才交人', '每次改動都跑；案例綁需求編號（v3 已提）'],
        ['上版／交付', '審查、上線', '審查紀錄、核准紀錄\nRelease Note、Deploy Guide', '④ 核准合併：程式負責人\n⑤ 授權上線：發布經理', 'AI 先做多輪審查並排嚴重度；rollback 先演練'],
        ['（新增）維運', '維運', '事件紀錄\n新的需求說明', '服務負責人', '異常由 AI 先診斷，結果回到需求分析']]
grid(s, 0.45, 1.4, [1.55, 1.25, 2.9, 3.35, 3.4], ['v3 階段', '手冊階段', '交接檔案（下一階段的輸入）', '人工關卡（誰核准）', '這一版補充'], ROWS, 0.74, size=10.5, hot=(5,))
note(s, f'①–⑤ 為手冊的五個人工關卡；寫程式的 AI 沒有核准權限。來源：{PB}，第 6、13、17、20、39、46 頁；{V3}。', y=6.62)

# ───────── 8 循環圖 ─────────
s = content('從直線到循環：v3 的核心原則不變，多了兩件事：關卡有人簽、上線後的問題回到需求')
ST = [('需求分析', '需求說明\nPrototype', '① 接受需求'), ('SA&D', '規格書\nERD、API', '② 確認規格'), ('開發', '實作計畫\n程式', '③ 核准計畫'),
      ('測試', '測試結果\n報告', 'AI 自行驗證'), ('上版／交付', '審查與\n核准紀錄', '④ 合併　⑤ 上線'), ('維運', '事件紀錄', '服務負責人')]
W, G, X0 = 1.85, 0.25, 0.5
for i, (h, f, g) in enumerate(ST):
    x = X0 + i * (W + G)
    tb(s, x, 1.5, W, 0.62, h, size=13, bold=True, color=WHITE, fill=OR if i < 5 else OR2, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    tb(s, x, 2.18, W, 0.78, f, size=10.5, fill=LG, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    tb(s, x, 3.02, W, 0.42, g, size=10.5, bold=True, color=OR, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    if i < 5: arrow(s, x + W + 0.03, 1.66, 0.19, 0.3)
# 回饋線：維運 → 需求分析
xl, xr = X0 + W / 2, X0 + 5 * (W + G) + W / 2
hline(s, xr, 3.44, xr, 3.95, OR, 2.25); hline(s, xl, 3.95, xr, 3.95, OR, 2.25); hline(s, xl, 3.62, xl, 3.95, OR, 2.25)
arrow(s, xl - 0.13, 3.46, 0.26, 0.22, MSO_SHAPE.UP_ARROW)
tb(s, 3.6, 3.78, 6.1, 0.34, '上線後的異常，由 AI 先診斷，結果寫成新的需求說明，回到第一步', size=10.5, bold=True, color=OR, fill=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 0.45, 4.35, 6.1, 2.15, [[('v3 原本就有，保留', {'bold': True, 'color': OR, 'size': 13})],
                              '・AI 先產出初稿，IT Review、修正、定稿後才進下一階段',
                              '・每階段有明確的 Input 與主要產出',
                              '・先 plan 再 code；限制修改範圍；不確定就標「待確認」'], size=12, fill=LG, margin=0.16)
tb(s, 6.8, 4.35, 6.1, 2.15, [[('這一版新增', {'bold': True, 'color': OR, 'size': 13})],
                             '・每個關卡有指定核准人，核准紀錄和產出放在一起',
                             '・交接靠檔案，不靠開會：下一階段直接讀上一階段的產出',
                             '・維運接回需求分析，同類問題不會一再發生'], size=12, fill=TINT, margin=0.16)
note(s, f'來源：{PB}，第 5 頁（從直線到循環）；{V3}，第 4 頁。', y=6.62)

# ───────── 9 需求分析分工 ─────────
s = content('需求分析分工建議：BA 講清楚業務、IT 轉成系統、AI 負責追問與起草')
ROWS = [['對焦業務問題', '主導：和業務單位釐清問題、使用者、成功標準', '參與：提早說出技術與系統限制', '追問範圍、使用者、限制，整理成需求說明初稿'],
        ['需求說明定稿', '撰寫與確認，送產品負責人', '檢視可行性，標出風險', '標出缺口與「待確認」，不自行補需求'],
        ['畫面與欄位', '確認符合業務流程與用語', '轉成系統／畫面規格（v3 IT Action ②）', '產生 Prototype、欄位清單、BPMN 草稿'],
        ['① 接受需求', '準備紀錄，說明取捨', '確認可以進入 SA&D', '不參與決定'],
        ['交給 SA&D', '回答業務問題', '主導 SA&D', '依需求說明起草規格']]
grid(s, 0.45, 1.4, [2.0, 3.55, 3.45, 3.45], ['步驟', 'BA', 'IT', 'AI'], ROWS, 0.72, size=11, hot=(3,))
callout(s, 5.5, '為什麼這樣分：', '需求說明寫得越清楚，後面四個階段的 AI 產出越準；這也是 v3 第 6 頁「Input 完整 → AI 產初稿」的前提。BA 熟悉保險業務，最適合當這一步的負責人。', h=0.85)
note(s, 'v3 第 6 頁標註「待與 BA 討論」。此為提案，請 Rex、Kenny 與 BA 在會議上確認；產品負責人由客戶或業務單位指定。', y=6.5)

# ───────── 10 矩陣 ─────────
s = content('哪些工作先交給 AI：用「多常發生 × 需要判斷」把 v3 列出的 AI 任務分成四類')
MX, MY, QW, QH = 1.15, 1.75, 3.6, 2.12
Q = [(0, 0, '常發生 × 判斷少', 'AI 先做，IT 確認', ['欄位清單', 'Test Case、Test Data', 'Unit Test、mock data', '註解與開發文件', 'Release Note'], TINT),
     (1, 0, '常發生 × 判斷多', 'AI 起草，IT 逐項審查', ['程式碼', 'API 規格（OpenAPI）', 'ERD、資料表', 'Code Review 風險檢查', '測試失敗原因分析'], LG),
     (0, 1, '少發生 × 判斷少', '做成範本，重複使用', ['Deploy Guide、User Guide', 'Migration Script', 'BPMN 圖的格式', '部署腳本、CI/CD'], LG),
     (1, 1, '少發生 × 判斷多', '人決定，AI 只提供資料', ['接受需求', '架構與技術選型', '核准合併、授權上線', '事件與例外處置'], WHITE)]
for c, r, h, rule, items, bg in Q:
    x, y = MX + c * (QW + 0.06), MY + r * (QH + 0.06)
    tb(s, x, y, QW, QH, '', fill=bg, line=LINE)
    tb(s, x + 0.12, y + 0.08, QW - 0.24, 0.62, [[(h, {'bold': True, 'size': 11, 'color': GREY})], [(rule, {'bold': True, 'size': 13, 'color': OR})]], margin=0.02)
    tb(s, x + 0.12, y + 0.74, QW - 0.24, QH - 0.82, '、'.join(items), size=11.5, margin=0.02)
tb(s, 0.45, MY, 0.65, QH, '常發生', size=11, bold=True, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
tb(s, 0.45, MY + QH + 0.06, 0.65, QH, '少發生', size=11, bold=True, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
tb(s, MX, MY + 2 * QH + 0.1, 2 * QW + 0.06, 0.3, '需要判斷：少 ─────────────→ 多', size=10.5, bold=True, color=GREY, align=PP_ALIGN.CENTER, margin=0.0)
tb(s, 8.75, 1.75, 4.15, 4.6, [[('怎麼用', {'bold': True, 'color': OR, 'size': 13})],
                              '・左上角最適合先做：最常發生、規則最清楚，最快看到效果，也最容易做成工具包。',
                              '・右上角量最大，但要人逐項審查；AI 讓審查變快，不能讓審查消失。',
                              '・右下角永遠由人決定，對應五個人工關卡。',
                              [('', {'size': 6})],
                              [('出處：', {'bold': True, 'color': OR}), (f'{SMR}：先把服務拆成活動，用「頻率」與「認知複雜度」評估，常發生、判斷少的活動最適合產品化。', {})]],
   size=11.5, fill=LG, margin=0.16)
note(s, f'任務清單取自 {V3} 各階段「AI 可以做什麼」；分類為提案，請 IT 團隊依實際專案調整。', y=6.62)

# ───────── 章節 03 ─────────
section('03', 'AI 協作\n怎麼落地', '怎麼讓 AI 每次都照規矩做，而不是看每個人 Prompt 寫得好不好？', ['從 Prompt 到範本、技能與自動閘門的四層做法', 'Review Checklist 升級版', '五頁案例頁怎麼填'])

# ───────── 12 四層控制 ─────────
s = content('從寫好 Prompt 到把規則寫進工具：v3 的 Prompt 六要素是第一層，工具包要往下做三層')
LV = [('1', 'Prompt 六要素', 'v3 附錄已有', '背景、權威輸入、角色、具體任務、限制、輸出格式。寫得好不好看個人。', '效果因人而異'),
      ('2', '範本與團隊說明檔', '這一版新增', '需求說明、規格、實作計畫用固定範本；團隊慣例寫進 CLAUDE.md，AI 每次都會先讀。', '格式一致、可以交接'),
      ('3', '技能（Skill）', '這一版新增', '資安、API、命名、法遵規範寫成檔案，AI 做事時自動套用；規範改了，下次使用就是新版。', '讓違規變少'),
      ('4', '自動閘門（Hook）與集中設定', '這一版新增', 'AI 每個動作前先檢查：沒有核准不能合併、不能上線；不能讀金鑰與機密檔。個人關不掉。', '讓違規幾乎做不到')]
for i, (n, h, tag, body, eff) in enumerate(LV):
    x = 0.45 + i * 3.14
    circle(s, x, 1.45, 0.55, n, fill=MID if i == 0 else OR, size=14)
    tb(s, x + 0.65, 1.45, 2.35, 0.55, tag, size=10.5, bold=True, color=GREY if i == 0 else OR, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x, 2.1, 3.0, 0.62, h, size=13, bold=True, color=WHITE, fill=MID if i == 0 else OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
    tb(s, x, 2.77, 3.0, 2.05, body, size=11.5, fill=LG, margin=0.14)
    tb(s, x, 4.87, 3.0, 0.48, eff, size=11.5, bold=True, color=OR, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    if i < 3: arrow(s, x + 3.0, 3.6, 0.14, 0.3)
callout(s, 5.6, '建議：', '工具包第一版做第 2 層（需求說明、規格、計畫三份範本）和一條第 3 層技能（例如「入口網站不能新增個資」），再用第 4 層守住「沒有核准不能上線」。', h=0.8)
note(s, f'手冊原話：技能讓違規變少，鉤子讓違規幾乎不可能。來源：{PB}，第 25–27、40–46 頁。', y=6.55)

# ───────── 13 Review Checklist ─────────
s = content('Review Checklist 升級：v3 的五項保留，再加四項保險業一定會被問的檢查')
ROWS = [['正確性', '是否符合需求？', 'IT Review', 'v3'],
        ['完整性', '有沒有遺漏情境（反向、邊界）？', 'IT Review；AI 先列缺口', 'v3'],
        ['一致性', '文件、API、程式是否一致？', 'IT Review；可用技能自動比對', 'v3'],
        ['安全性', '權限、輸入驗證、敏感資料？', 'IT Review；AI 先做資安初審', 'v3'],
        ['可追溯性', '產出能不能回到需求來源？', '產出標需求編號', 'v3'],
        ['法遵', '有沒有用到個資？資料是否在可給 AI 讀的範圍內？', '技能套用規範；資料分級由 IT 設定', '新增'],
        ['職責分離', '寫程式的 AI 有沒有自己核准？關卡是否由指定的人簽？', '自動閘門強制', '新增'],
        ['測試證據', '有沒有自動測試結果？案例是否綁需求編號？', 'CI 自動產生', '新增'],
        ['紀錄完整', '誰核准、何時、套用哪個版本的規範，查得到嗎？', '版本紀錄與核准紀錄', '新增']]
grid(s, 0.45, 1.4, [1.6, 5.4, 4.15, 1.3], ['檢查項目', '問什麼', '誰檢查、怎麼檢查', '來源'], ROWS, 0.5, size=11, hot=(5, 6, 7, 8))
note(s, f'淺橘底為新增項目，對應金管會 AI 指引的治理與問責、透明可解釋原則。來源：{V3}，第 16 頁；{PB}。', y=6.4)

# ───────── 14 案例頁 ─────────
s = content('五頁案例頁怎麼填：用同一個案例走完五個階段，每階段截三張圖、記一個數字')
ROWS = [['需求分析', '業務需求原文、畫面需求', 'AI 追問的問題、需求說明、Prototype', '修正前後差異', '到需求定稿花多久'],
        ['SA&D', '需求說明、既有資料表', 'ERD、API 規格、Test Plan', '技術 Review 意見、標「待確認」的欄位', '開發後還要改規格幾次'],
        ['開發', '規格、既有程式、Coding Convention', '實作計畫、程式差異、Unit Test', 'Code Review 意見與修改', '第一次就能合併的比例'],
        ['測試', '驗收條件', '正向、反向、邊界案例；失敗原因分析', '補了哪些案例', '第一次通過 CI 的比例'],
        ['上版／交付', '版本差異、測試結果', 'Release Note、Deploy Guide 草稿', '核准紀錄、rollback 計畫', '上線失敗率']]
grid(s, 0.45, 1.4, [1.55, 2.75, 3.3, 2.85, 2.0], ['階段', '① IT Input 截什麼', '② AI 產出截什麼', '③ IT Review 截什麼', '記錄的數字'], ROWS, 0.62, size=10.5)
tb(s, 0.45, 4.95, 6.1, 1.25, [[('案例怎麼挑', {'bold': True, 'color': OR, 'size': 12})],
                              '優先用 IT 團隊已經做過的真實專案；沒有的話，用「理賠進度自助查詢」做示範專案。範圍小、規格清楚、已有測試最好。'],
   size=11.5, fill=LG, margin=0.14)
tb(s, 6.8, 4.95, 6.1, 1.25, [[('一次錄、三處用', {'bold': True, 'color': OR, 'size': 12})],
                             '同一批畫面同時用在：這五頁案例頁、第 3 集影片第③段示範（D1–D4）、對客簡報「為什麼是 PwC FSC1」的實績欄。'],
   size=11.5, fill=TINT, margin=0.14)
note(s, '截圖與錄影不可含客戶資料或真實個資；示範資料一律用假資料。', y=6.35)

# ───────── 章節 04 ─────────
section('04', '治理、法規\n與成效', '保險業客戶會問哪些法規問題？怎麼證明這套流程有效？', ['保險業法規與所內規定', '每個階段要量的指標'])

# ───────── 16 法規 ─────────
s = content('法規與所內規定：用外部 AI 處理程式與需求可能涉及委外，導入前就要一起規劃')
ROWS = [['《保險業作業委託他人處理內部作業制度及程序辦法》', '2025-12-31 施行', '程式碼與需求文件交給外部 AI 處理，可能屬作業委外，要判斷是否申報、納入委外管理'],
        ['《保險業作業委外使用雲端服務自律規範》', '壽險、產險公會訂定', 'AI 服務架在雲端時，雲端業者的管理與查核安排要一併處理'],
        ['《保險業運用人工智慧系統自律規範》', '2025-04-10 金管會備查', '保險公司對 AI 結果負最終責任；AI 工具要登錄 AI 清冊、納入風險管理'],
        ['金管會《金融業運用 AI 指引》', '2024-06-20 發布', '治理與問責、透明可解釋、穩健與安全等原則，都要有對應做法與紀錄']]
grid(s, 0.45, 1.4, [4.3, 2.2, 5.95], ['法規或規範', '現況', '對這套流程的意思'], ROWS, 0.72, size=11)
tb(s, 0.45, 4.6, 12.45, 0.42, '所內規定（我們自己做的時候也要遵守）', size=12.5, bold=True, color=WHITE, fill=INK, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
RULES = [('客戶資料', '不放進外部 AI 工具；示範一律用假資料'), ('AI 產出', '對外使用前一定要人工確認'), ('獨立性', '客戶若同時是審計客戶，推案前先和 Risk & Quality 確認')]
for i, (h, t) in enumerate(RULES):
    tb(s, 0.45 + i * 4.18, 5.08, 4.05, 0.95, [[(h, {'bold': True, 'color': OR, 'size': 12})], t], size=11.5, fill=LG, margin=0.12)
note(s, '是否屬委外、是否需申報，須依客戶個案與實際架構判斷並與客戶法遵確認；本頁非法律意見。日期依公開資料整理。', y=6.2)

# ───────── 17 成效指標 ─────────
s = content('怎麼證明有效：每個階段一個先看的指標，數據大多已經在版本紀錄和 CI 裡')
ROWS = [['需求分析', '從第一次討論到需求說明定稿的時間', '被接受、進入 SA&D 的需求比例'],
        ['SA&D', '需求說明到規格定稿的時間', '開發開始後還要改規格的次數'],
        ['開發', '第一次實作就能合併的比例', '每個改動的返工次數'],
        ['測試', 'AI 改動第一次就通過 CI 的比例', '上線前抓到的缺陷 vs 上線後才發現的'],
        ['上版／交付', '第一次審查的等待時間', '上線失敗率、復原時間'],
        ['維運', '異常發生到寫成需求說明的時間', '同類事件重複發生的次數']]
grid(s, 0.45, 1.4, [1.6, 5.6, 5.25], ['階段', '先看（領先指標）', '後看（落後指標）'], ROWS, 0.6, size=11.5)
callout(s, 5.55, '建議：', '試點前先記下現況基準，試點後用同一套指標比較。手冊說的「數週縮短到數小時」是預期，對外只講我們自己量到的數字。', h=0.75)
note(s, f'整體交付績效可再搭配 DORA 指標（部署頻率、交付時間、變更失敗率、復原時間）。來源：{PB}，各階段「How to measure it」。', y=6.45)

# ───────── 章節 05 ─────────
section('05', '對外：服務化\n與影片', '這套流程怎麼變成可以賣的服務？影片怎麼講？', ['你自己做、陪你做、幫你做三種合作方式', '同業已經在做什麼、我們的差異', '第 3 集影片和這份簡報的對應'])

# ───────── 19 三種合作方式 ─────────
s = content('三種合作方式：同一套流程，依客戶想要多少人力參與，包成三種方案')
TIER = [('你自己做', '工具包', ['需求說明、規格、計畫三份範本', '保險業規範技能（資安、個資、委外）', '自動閘門設定與 Review Checklist'], '已有 IT 團隊、只缺方法的客戶', MID),
        ('陪你做', '試點陪跑＋教學', ['現況診斷，挑一個試點', '帶團隊用真實需求走完五個階段', '各角色教學、試點前後成效比較'], '多數客戶的起點', OR),
        ('幫你做', '駐點開發', ['FSC1 人力駐點，用這套流程開發', '交付程式、文件與完整紀錄', '期末交接工具包與做法'], '人力不足、時程緊的客戶', OR2)]
for i, (h, sub, items, fit, c) in enumerate(TIER):
    x = 0.45 + i * 4.18
    tb(s, x, 1.42, 4.05, 0.75, [[(h, {'bold': True, 'size': 16}), ('　' + sub, {'size': 12})]], color=WHITE, fill=c, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
    tb(s, x, 2.22, 4.05, 2.3, [[('・' + t, {})] for t in items], size=12, fill=LG, margin=0.16)
    tb(s, x, 4.57, 4.05, 0.62, [[('適合：', {'bold': True, 'color': OR}), (fit, {})]], size=11.5, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
callout(s, 5.45, '建議主推：', '「陪你做一個試點」。風險低、客戶看得到成效，做完自然延伸到工具包或駐點。對應 SDLC 完整版的四個服務模組：① 診斷、② 規範與控制、③ 陪跑教學、④ 量測擴大。', h=0.85)
note(s, f'三種方案的概念出自 {SMR}（Do it for me／with me／yourself）。範圍與報價待 Sam、Luke 決定，影片不提價格。', y=6.42)

# ───────── 20 同業 ─────────
s = content('同業已經在做：保險顧問業把服務做成產品很普遍，我們要講清楚差在哪裡（內部參考）')
ROWS = [['Deloitte', 'Agentic AI Blueprint for Insurance', '產險／壽險 AI 導入方法，從評估、設計到試點、規模化，含預建 AI 代理', '最接近我們的 AI 協作開發服務'],
        ['Capgemini', 'Guidewire 生成式 AI／代理工具組', '舊商品遷移、需求分析、系統整合的 AI 加速器', '和第 3 集「核心系統」主題直接重疊'],
        ['Deloitte、KPMG', 'IFRS 17 加速器（與 SAS 合作）', '預建資料模型與計算引擎，強調數週到 90 天導入', '客戶做 IFRS 17 時可能拿來比較'],
        ['McKinsey、BCG', '核心現代化、AI-First 藍圖', '方法論與藍圖，尚未做成可單獨購買的產品', '代表「框架」層級，還不是產品']]
grid(s, 0.45, 1.4, [1.85, 3.1, 4.6, 2.9], ['公司', '方案', '做什麼', '對我們的意思'], ROWS, 0.72, size=10.5, hot=(0, 1))
tb(s, 0.45, 4.75, 12.45, 0.42, '我們的差異（影片與對客簡報都用這三點，不點名同業）', size=12.5, bold=True, color=WHITE, fill=INK, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
DIFF = [('台灣保險法規', '委外辦法、公會 AI 與雲端自律規範直接寫進技能與檢查'), ('在地核心系統', '熟悉理賠、核保、核心系統與 IFRS 17，試點直接用客戶的業務情境'), ('可稽核的人工關卡', '五個關卡誰核准、什麼時候核准，全部查得到')]
for i, (h, t) in enumerate(DIFF):
    tb(s, 0.45 + i * 4.18, 5.22, 4.05, 1.0, [[(h, {'bold': True, 'color': OR, 'size': 12})], t], size=11.5, fill=TINT, margin=0.12)
note(s, '資料來自業者新聞稿與產品頁，未經獨立驗證；連結見 research/productization/保險顧問產品化服務盤點.md。對外資料不點名公司。', y=6.35)

# ───────── 21 影片對應 ─────────
s = content('第 3 集影片〈核心系統，跟得上 AI 嗎？〉：每一段都對應這份簡報，示範要靠案例錄影')
ROWS = [['開場', '一個問題帶出主題', '—', '影片組'],
        ['① 導入前後對照＋矩陣', '直線變循環；「常常發生、規則清楚的交給 AI，少見、需要判斷的留給人」', '第 8、10 頁', '影片組'],
        ['② 需求範本', 'AI 用固定範本把需求整理成格式一致的文件', '第 12 頁第 2 層', 'Rex、BA（範本內容）'],
        ['③ 預錄示範', 'AI 在各階段怎麼用，最後由人決定能不能上線', '第 14 頁案例錄影', 'Rex、Kenny（錄影 D1–D4）'],
        ['④ 工具包＋三種合作方式', '工具包依台灣法規與核心系統調整；你自己做、陪你做、幫你做', '第 19、20 頁', 'Sam、Luke（範圍與 CTA）']]
grid(s, 0.45, 1.4, [2.6, 5.4, 1.95, 2.5], ['影片段落', '畫面與旁白重點', '對應本簡報', '需要誰提供'], ROWS, 0.7, size=11, hot=(3,))
callout(s, 5.4, '時程：', '腳本 v0.3（22 個鏡頭、約 128 秒）已完成；10/28 與 Sam 檢視腳本；11/04 獨立性檢查。示範錄影是目前最大的缺口。', h=0.75)
note(s, '腳本：pwc-fsc1-ai-short-vid/08-core-system-ai/script.md。影片不點名公司、不使用真人肖像。', y=6.35)

# ───────── 章節 06 ─────────
section('06', '下一步', '今天會議要決定什麼？誰負責哪一件？', ['四個要決定的問題', '分工與期限'])

# ───────── 23 下一步 ─────────
s = content('今天要決定的四件事，以及會後分工')
DEC = [('案例', '五頁案例頁用哪一個專案？IT 團隊的真實專案，還是理賠進度自助查詢示範？'),
       ('需求分析分工', '第 9 頁的 BA × IT × AI 分工是否採用？'),
       ('工具包第一版', '先做三份範本＋一條技能＋上線閘門，是否同意？'),
       ('錄影', '誰錄、用哪個環境、什麼時候錄（同時給影片 D1–D4）？')]
tb(s, 0.45, 1.4, 4.6, 0.42, '要決定的事', size=12.5, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
for i, (h, t) in enumerate(DEC):
    y = 1.88 + i * 1.08
    circle(s, 0.45, y + 0.25, 0.45, str(i + 1), size=12)
    tb(s, 1.0, y, 4.05, 1.0, [[(h, {'bold': True, 'size': 12})], t], size=10.5, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
ROWS = [['選定案例並錄製五階段畫面', 'Rex、Kenny', '會議決定'],
        ['填寫 IT 團隊實績填寫表', 'Kenny／IT 團隊', '會議決定'],
        ['需求說明、規格、計畫三份範本初版', 'BA、Rex', '會議決定'],
        ['保險業規範技能與 Review Checklist 定稿', 'BA 起草、IT 確認', '會議決定'],
        ['工具包名稱與範圍、三種方案報價', 'Sam、Luke', '會議決定'],
        ['影片腳本檢視', '影片組、Sam', '10/28'],
        ['獨立性檢查', '影片組、R&Q', '11/04']]
grid(s, 5.35, 1.4, [3.9, 2.3, 1.35], ['事項', '負責', '期限'], ROWS, 0.6, size=10.5, first_bold=False, hot=(0,))
note(s, '負責人與期限為提案，以會議決定為準。', y=6.35)

# ───────── 24 來源 ─────────
s = content('附錄｜來源')
SRCS = [('上週版本', f'{V3}（19 頁）。本簡報的五階段架構、IT／AI 分工、Review Checklist 與 Prompt 六要素皆沿用此版。'),
        ('流程與控制', f'Louis Claxton，{PB}，2026。人工關卡、交接檔案、技能與鉤子、成效指標。完整白話說明見 research/ai-native-sdlc/AI原生SDLC_白話說明.pptx。'),
        ('產品化框架', f'Mohanbir Sawhney，〈How to Turn Professional Services Into Products〉，MIT Sloan Management Review，2024。中文整理見 research/productization/MIT-SMR_專業服務產品化_文章報告.md。'),
        ('同業盤點', 'research/productization/保險顧問產品化服務盤點.md（業者新聞稿與產品頁，含連結）。'),
        ('法規', '保險業作業委託他人處理內部作業制度及程序辦法；保險業作業委外使用雲端服務自律規範；保險業運用人工智慧系統自律規範；金管會《金融業運用人工智慧（AI）指引》。'),
        ('影片', '第 3 集腳本 v0.3：pwc-fsc1-ai-short-vid/08-core-system-ai/script.md。')]
for i, (h, t) in enumerate(SRCS):
    y = 1.4 + i * 0.83
    tb(s, 0.45, y, 2.0, 0.75, h, size=12, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
    tb(s, 2.5, y, 10.4, 0.75, t, size=11, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.12)

p.save(OUT)
print(OUT, len(p.slides), 'slides')
