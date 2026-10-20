# 檢查簡報文字是否可能超出文字框、表格是否撐出頁面（估算值，偏保守）
# 用法：python3 check_fit.py <簡報.pptx>
import sys, math, unicodedata
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import PP_PLACEHOLDER

EMU = 914400
SAFE_BOTTOM = 7.05  # 頁尾上緣

def char_w(ch):
    if unicodedata.east_asian_width(ch) in ('W', 'F'): return 1.0
    if ch in 'ilI.,:;|!\'': return 0.3
    if ch.isupper() or ch in 'mwMW%@': return 0.72
    return 0.56

def text_height(tf, width_in, default_size):
    h = 0
    for para in tf.paragraphs:
        runs = para.runs
        if not runs: h += default_size * 1.2 / 72; continue
        size = max((r.font.size.pt if r.font.size else default_size) for r in runs)
        line_w = sum(char_w(c) * (r.font.size.pt if r.font.size else default_size) / 72 for r in runs for c in r.text)
        lines = max(1, math.ceil(line_w / max(width_in, 0.1) - 1e-6))
        sa = para.space_after.pt if para.space_after is not None else 0
        h += lines * size * 1.2 / 72 + sa / 72
    return h

def margins(tf, default=(0.1, 0.1, 0.05, 0.05)):
    m = [tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom]
    return [x / EMU if x is not None else d for x, d in zip(m, default)]

def check(path):
    p = Presentation(path)
    sw, sh = p.slide_width / EMU, p.slide_height / EMU
    issues = 0
    for n, s in enumerate(p.slides, 1):
        for shp in s.shapes:
            x, y = (shp.left or 0) / EMU, (shp.top or 0) / EMU
            w, h = (shp.width or 0) / EMU, (shp.height or 0) / EMU
            if x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + h > sh + 0.01:
                print(f'p{n}: {shp.name} 超出頁面 ({x:.2f},{y:.2f},{w:.2f},{h:.2f})'); issues += 1
            if shp.is_placeholder and shp.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                continue  # 頁尾與頁碼沿用範本字級
            if shp.has_text_frame and shp.text_frame.text.strip():
                tf = shp.text_frame
                ml, mr, mt, mb = margins(tf)
                need = text_height(tf, w - ml - mr, 18)
                if need > h - mt - mb + 0.04:
                    print(f'p{n}: 文字可能溢出 需 {need:.2f} 有 {h - mt - mb:.2f}｜{tf.text[:30]!r}'); issues += 1
            if shp.has_table:
                tbl = shp.table; total = 0
                for r_i, row in enumerate(tbl.rows):
                    need_row = row.height / EMU
                    for c_i, cell in enumerate(row.cells):
                        cw = tbl.columns[c_i].width / EMU
                        ml, mr, mt, mb = margins(cell, (0.1, 0.1, 0.05, 0.05))
                        need_row = max(need_row, text_height(cell.text_frame, cw - ml - mr, 18) + mt + mb)
                    total += need_row
                if y + total > SAFE_BOTTOM:
                    print(f'p{n}: 表格撐高後底部到 {y + total:.2f} 吋（上限 {SAFE_BOTTOM}）'); issues += 1
                elif total > h + 0.05:
                    print(f'p{n}: 表格會撐高 {h:.2f} → {total:.2f} 吋，確認不壓到下方元件')
    print(path, '問題數', issues)

for f in sys.argv[1:]:
    check(f)
