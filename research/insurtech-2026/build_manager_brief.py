# 保險科技服務優先順序：主管三頁精簡版（附可點擊來源連結，以所內簡報範本建立）
# 用法：python3 build_manager_brief.py <範本.pptx> <輸出.pptx>
import sys, copy, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TPL, OUT = sys.argv[1], sys.argv[2]
OR = RGBColor(0xFD, 0x51, 0x08); OR2 = RGBColor(0xFE, 0x7C, 0x39); TINT = RGBColor(0xFF, 0xE8, 0xD4)
# 範本佈景主題色：tx1 黑、lt2 淺灰、accent6 線條灰
INK = RGBColor(0x00, 0x00, 0x00); GREY = RGBColor(0x59, 0x59, 0x59); LG = RGBColor(0xEB, 0xEB, 0xEB); LINE = RGBColor(0xCB, 0xD1, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# 範本字型：標題 Georgia、內文 Arial、中文一律微軟正黑體（與所內範本範例頁一致）
FONT = 'Arial'; HEAD_FONT = 'Georgia'; EA_FONT = '微軟正黑體'

SYMBOLS = set('●◐○✓✗→①②③④⑤⚠↑')

def style_font(r, head=False):
    # Arial 沒有 ●◐○ 等符號，這類文字改用中文字型顯示
    r.font.name = EA_FONT if SYMBOLS & set(r.text) else (HEAD_FONT if head else FONT)
    rPr = r._r.get_or_add_rPr()
    for e in rPr.findall(qn('a:ea')): rPr.remove(e)
    rPr.find(qn('a:latin')).addnext(rPr.makeelement(qn('a:ea'), {'typeface': EA_FONT}))
FOOT = '內部討論稿｜僅供 FSC1 保險組內部使用｜資料來源為公開研究，點選來源可開啟原文'

# 來源連結：清單（insurtech_research_links_directory.md）有具體頁面者用清單連結；
# 清單沒有、或只給入口首頁者，用網路搜尋找到的具體頁面
URL = {
    'gallagher': 'https://www.ajg.com/gallagherre/news-and-insights/global-insurtech-report-q1-2026/',
    'gallagher_news': 'https://www.insurancebusinessmag.com/reinsurance/news/breaking-news/ai-startups-capture-record-95-2-of-insurtech-funding--gallagher-re-574376.aspx',
    'pwc57': 'https://amp.claimsjournal.com/news/national/2025/11/17/334094.htm',
    'bain': 'https://www.bain.com/ko/results/reimagining-insurance-for-the-ai-era',
    'mckinsey': 'https://www.mckinsey.com/industries/financial-services/our-insights/can-agentic-ai-finally-modernize-core-technologies-in-insurance',
    'deloitte': 'https://www.deloitte.com/ap/en/about/press-room/deloitte-launches-new-paper-on-scaling-agentic-ai-in-life-insurance.html',
    'fsc_ai': 'https://www.fsc.gov.tw/websitedowndoc?file=chfsc/202408231741530.pdf',
    'fsc_ai_news': 'https://www.leeandli.com/TW/Newsletters/7287.htm',
    'fsc_outsource': 'https://www.fsc.gov.tw/ch/home.jsp?id=96&parentpath=0,2&mcustomize=news_view.jsp&dataserno=202408200001',
    'outsource_rule': 'https://www.leeandli.com/TW/Newsletters/7525.htm',
    'cloud_selfreg': 'https://law.lia-roc.org.tw/Law/GetFile/0000419987',
    'pwc_ifrs17': 'https://www.pwc.tw/zh/topics/trends/ifrs-17-key-issues.html',
    'naic': 'https://content.naic.org/cmte_h_bdaiwg.htm',
    'naic_pilot': 'https://www.fenwick.com/insights/publications/naic-expands-ai-systems-evaluation-tool-pilot-program-to-12-states-key-updates-for-insurers-and-ai-vendors-supporting-insurers',
    'mas': 'https://www.mas.gov.sg/regulation/guidelines/guidelines-on-artificial-intelligence-risk-management-for-financial-institutions',
    # 競爭者（官網或原文頁）
    'mck_uwos': 'https://www.mckinsey.com/industries/financial-services/our-insights/insurance-blog/the-future-underwriting-operating-system-from-inbox-to-ai-nerve-center',
    'bcg_ap': 'https://www.bcg.com/assets/2026/white-paper-competitive-advantage-in-the-agentic-ai-era-ap-insurers.pdf',
    'bcg_pc': 'https://www.bcg.com/assets/2026/executive-perspectives-ai-first-companies-win-the-future-property-casualty-insurance.pdf',
    'bain_dist': 'https://www.bain.com/it/about-bain/media-center/press-releases/2024/$50-billion-opportunity-emerges-for-insurers-worldwide-from-generative-ais-potential-to-boost-revenues-and-take-out-costs/',
    'dtt_blueprint': 'https://www.deloitte.com/global/en/what-we-do/capabilities/agentic-ai-blueprint-for-insurance.html',
    'dtt_ap': 'https://www.deloitte.com/ap/en/perspectives/agentic-ai-scale-for-life-insurers-in-apac.html',
    'ey_nexus': 'https://www.ey.com/en_us/insurance/nexus',
    'kpmg_ii': 'https://kpmg.com/au/en/insights/artificial-intelligence-ai/intelligent-insurance.html',
    'kpmg_trust': 'https://kpmg.com/ca/en/insights/2024/11/you-can-build-trust-with-ai.html',
    'pwc_agentos': 'https://www.pwc.nl/en/services/artificial-intelligence/agent-os.html',
}

p = Presentation(TPL)
lst = p.slides._sldIdLst
for sid in list(lst):
    p.part.drop_rel(sid.rId); lst.remove(sid)
L = {l.name: l for l in p.slide_masters[2].slide_layouts}
# 移除範本遺留的投影片章節（PowerPoint 左側的章節名稱）
for ext in list(p.part._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}ext')):
    if ext.find('{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst') is not None:
        ext.getparent().remove(ext)

def set_ph(s, idx, text, size=None, bold=None):
    for ph in s.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            for para in ph.text_frame.paragraphs:
                for r in para.runs:
                    style_font(r, head=idx == 0)
                    if size: r.font.size = Pt(size)
                    if bold is not None: r.font.bold = bold
            return ph

def tb(s, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None, margin=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(0.75)
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'): setattr(tf, f'margin_{side}', Inches(margin))
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg in segs:
            t, st = (seg, {}) if isinstance(seg, str) else seg
            r = para.add_run(); r.text = t
            style_font(r); f = r.font; f.size = Pt(st.get('size', size)); f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color)
            if 'link' in st:
                r.hyperlink.address = st['link']; f.underline = True
        para.space_after = Pt(2)
    return sh

def circle(s, x, y, d, text, fill=OR, size=12):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    r = para.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; style_font(r)

def hline(s, x1, y1, x2, y2, color=LINE, width=1.0):
    ln = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(width)

def footer(s):
    # python-pptx 新增投影片不會帶入頁尾與頁碼，從版面複製（頁碼為自動欄位）
    have = {ph.placeholder_format.idx for ph in s.placeholders}
    for ph in s.slide_layout.placeholders:
        if ph.placeholder_format.idx in (3, 4) and ph.placeholder_format.idx not in have:
            s.shapes._spTree.append(copy.deepcopy(ph._element))
    set_ph(s, 3, FOOT)

def content(title):
    s = p.slides.add_slide(L['只有標題'])
    set_ph(s, 0, title, 20, bold=True); footer(s)
    return s

def sources(s, items):
    segs = [('來源：', {'bold': True, 'color': GREY})]
    for i, (label, key) in enumerate(items):
        if i: segs.append(('　', {}))
        segs.append((label, {'link': URL[key]}))
    tb(s, 0.45, 6.62, 12.45, 0.42, [segs], size=9, color=GREY, margin=0.03)

# ───────── 第 1 頁 重點摘要 ─────────
s = content('重點摘要：保險 AI 已進入規模化，建議第一波先推 IFRS 17 與 AI 治理兩項服務包')
STATS = [('95.2%', '2026 Q1 全球保險科技募資流向 AI 企業的比例（歷史新高）', 'Gallagher Re, 2026-05', 'gallagher'),
         ('57%', '保險高階主管把生成式與代理型 AI 列為 2026 首要科技投資', 'PwC 調查, 2025-11', 'pwc57'),
         ('95%', '投資 AI 但尚未看到可衡量回報的企業（跨產業）', 'Bain 引述 MIT', 'bain')]
for i, (n, d, src, key) in enumerate(STATS):
    y = 1.4 + i * 1.72
    hot = i == 2
    tb(s, 0.45, y, 3.9, 1.58, '', fill=OR if hot else LG)
    tb(s, 0.6, y + 0.06, 3.6, 0.65, n, size=32, color=WHITE if hot else OR)
    tb(s, 0.6, y + 0.72, 3.6, 0.5, d, size=11, color=WHITE if hot else INK, margin=0.03)
    tb(s, 0.6, y + 1.24, 3.6, 0.3, [[(src, {'link': URL[key], 'color': WHITE if hot else INK})]], size=9, margin=0.03)
tb(s, 4.6, 1.4, 8.3, 1.1, [[('情境　', {'bold': True, 'color': OR}), ('資金與預算全面轉向 AI，保險業已從試點走向規模化。', {})],
                           [('衝突　', {'bold': True, 'color': OR}), ('多數投資還看不到回報；台灣同時面對 IFRS 17／TW-ICS 首年出表與 AI 治理要求。', {})]],
   size=12, fill=LG, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
tb(s, 4.6, 2.62, 8.3, 0.9, [[('建議　', {'bold': True, 'size': 12}), ('第一波先推 ① IFRS 17／TW-ICS 後財務轉型　② AI 清冊、模型驗證與護欄；代理型 AI 核保理賠、核心現代化、深偽反詐欺列為第二波。', {})]],
   size=13.5, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.16)
WHY = [('需求有時限', 'IFRS 17／TW-ICS 2026-01-01 接軌；委外新辦法 2025-12-31 施行；AI 指引進入落地'),
       ('離既有服務最近', '延伸既有駐點團隊與 AI 治理服務，不需新技術夥伴；AI 治理健檢 4–6 週可交付'),
       ('能帶出第二波', '財務資料與 AI 治理是代理型 AI、核心現代化的前置條件，可延伸為 1–3 年人力委外')]
for i, (h, d) in enumerate(WHY):
    x = 4.6 + i * 2.8
    tb(s, x, 3.65, 2.68, 0.48, h, size=12.5, bold=True, color=OR, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    tb(s, x, 4.18, 2.68, 1.4, d, size=11, fill=LG, margin=0.12)
tb(s, 4.6, 5.72, 8.3, 0.75, [[('今天請決定：', {'bold': True, 'color': OR}), ('同意第一波兩項服務，並各指派 1 位負責人。', {})]],
   size=12, fill=TINT, anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
sources(s, [('Gallagher Re Global InsurTech Report Q1 2026', 'gallagher'), ('Insurance Business 報導（95.2%）', 'gallagher_news'),
            ('PwC 保險高階主管調查（Claims Journal）', 'pwc57'), ('Bain：Reimagining Insurance for the AI Era', 'bain')])

# ───────── 文獻清單（第 2 頁引用編號、附錄頁與 文獻清單.md 共用） ─────────
# (類別, 標題, 機構與日期, 網址, 網址已由網路搜尋確認)
REFS = [
    ('市場數據', 'Global InsurTech Report Q1 2026', 'Gallagher Re，2026-05', URL['gallagher'], False),
    ('市場數據', 'AI start-ups capture record 95.2% of insurtech funding', 'Insurance Business，2026-05', URL['gallagher_news'], True),
    ('市場數據', 'PwC：保險高階主管視代理型 AI 為轉型主力（57%）', 'Claims Journal，2025-11-17', URL['pwc57'], True),
    ('市場數據', 'Reimagining Insurance for the AI Era', 'Bain（引述 MIT 2025）', URL['bain'], True),
    ('IFRS 17／TW-ICS', '保險局全球資訊網（IFRS 17／TW-ICS 專區）', '金管會保險局', 'https://www.ib.gov.tw/', False),
    ('IFRS 17／TW-ICS', '保險業資本新時代：壽險公會談 15 年過渡措施', '優分析，2025-10', 'https://uanalyze.com.tw/articles/4909935653', True),
    ('IFRS 17／TW-ICS', 'Insurance Banana Skins 2025：台灣焦慮高於全球、準備度低於全球', 'PwC／CSFI（中央社轉載），2025-10-07', 'https://n.yam.com/Article/20251007545281', True),
    ('IFRS 17／TW-ICS', 'IFRS 17 即將上路：建議臺灣保險業者的關鍵議題', '資誠，2025-11', URL['pwc_ifrs17'], False),
    ('AI 治理', '《金融業運用 AI 指引》', '金管會，2024-06', URL['fsc_ai'], False),
    ('AI 治理', '金融業運用 AI 指引解析', '理律法律事務所', URL['fsc_ai_news'], True),
    ('AI 治理', '《保險業運用人工智慧系統自律規範》', '壽險／產險公會，2025-04-10 備查', 'https://law.lia-roc.org.tw/Law/GetFile/0000394231', True),
    ('AI 治理', 'NAIC expands AI Systems Evaluation Tool pilot to 12 states', 'Fenwick，2026', URL['naic_pilot'], True),
    ('AI 治理', 'Guidelines on AI Risk Management for Financial Institutions', 'MAS，2026-10 定稿', URL['mas'], True),
    ('代理型 AI', 'Shedding Light on Agentic AI in Insurance（22% 計畫 2026 年底前上線）', 'Celent，2025-05', 'https://www.celent.com/en/insights/shedding-light-on-agentic-ai-in-insurance', True),
    ('代理型 AI', 'Manulife MAUDE 獲 Celent 2026 Model Insurer（核保 AI）', 'Manulife，2026', 'https://www.manulife.com/ph/about-us/news/manulife-earns-celent-model-insurer-recognition-for-ai-leadership-in-underwriting', True),
    ('代理型 AI', 'The future underwriting operating system（UWOS）', 'McKinsey，2026-06', URL['mck_uwos'], True),
    ('核心與雲端', 'Can agentic AI (finally) modernize core technologies in insurance?', 'McKinsey，2026-04', URL['mckinsey'], True),
    ('核心與雲端', '保險業作業委外新辦法（2025-12-31 施行）', '理律法律事務所', URL['outsource_rule'], True),
    ('核心與雲端', '《保險業作業委外使用雲端服務自律規範》', '壽險／產險公會，2026-06 修正', URL['cloud_selfreg'], True),
    ('深偽反詐欺', 'SONAR 2025：How deepfakes, disinformation and AI amplify insurance fraud', 'Swiss Re Institute，2025-06', 'https://www.swissre.com/institute/research/sonar/sonar2025/how-deepfakes-disinformation-ai-amplify-insurance-fraud.html', True),
    ('深偽反詐欺', 'Deepfakes put insurers in deep water', 'CAS Actuarial Review，2026-05／06', 'https://digital.casact.org/issue/may-june-2026/deepfakes-put-insurers-in-deep-water/', True),
    ('深偽反詐欺', 'Deepfakes: A growing threat to digital media integrity for insurers', 'Verisk', 'https://www.verisk.com/resources/campaigns/deepfakes-a-growing-threat-to-digital-media-integrity-for-insurers/', True),
]

def cites(nums):
    segs = []
    for n in nums:
        segs.append((f'[{n}]', {'link': REFS[n - 1][3], 'size': 10}))
    return [segs]

# ───────── 第 2 頁 五項服務的依據 ─────────
s = content('五項優先服務的依據：第一波兩項同時具備監理時限、既有能力與 4–8 週切入案')
F, M, N = '●', '◐', '○'
SVC = [('IFRS 17／TW-ICS 後財務轉型',
        '2026-01-01 與 TW-ICS 同步上路，15 年過渡期逐年趨緊；台灣業者對新制的準備度低於全球（3.06 對 3.27）',
        '新制以現時價值衡量負債，業者經營指標轉向 CSM 與資產負債管理（ALM）',
        '首年結帳瓶頸診斷＋CSM 管理儀表板', (F, F, F), '第一波', [5, 6, 7, 8]),
       ('AI 清冊、模型驗證與護欄',
        '金管會 AI 指引（2024-06）；公會 AI 自律規範 2025-04-10 備查，保險公司對 AI 結果負最終責任',
        'NAIC 12 州 AI 評估工具試點（2026-03～09）；MAS AI 風險管理指引 2026-10 定稿',
        '4–6 週 AI 清冊盤點＋自律規範合規健檢', (F, F, F), '第一波', [9, 10, 11, 12, 13]),
       ('代理型 AI 核保與理賠',
        '多數停在單點智能助理；Celent：22% 保險公司計畫 2026 年底前上線代理型 AI',
        'Manulife MAUDE 自動核准逾 58% 合格壽險案件、約 2 分鐘通知；McKinsey 核保作業系統',
        '6–8 週單一險種工作流藍圖', (F, M, M), '第二波', [14, 15, 16]),
       ('AI 加速核心現代化與雲端',
        '委外新辦法 2025-12-31 施行、雲端自律規範 2026-06 修正；舊核心難以支援 API 與 AI 串接',
        'McKinsey：代理型 AI 可壓縮舊程式解析、測試與對帳的重工',
        '雲端委外合規評估，或周邊系統 AI 重寫 PoC', (M, M, N), '第二波', [17, 18, 19]),
       ('深偽理賠反詐欺',
        '修圖車損、偽造醫療單據門檻降低，規則引擎難以辨識',
        'Swiss Re SONAR 2025 將「深偽助長保險詐欺」列為高影響新興風險',
        '車損照片／醫療單據偽造偵測 PoC', (M, N, M), '第二波', [20, 21, 22])]
HEAD = ['服務', '為什麼是現在（監理與痛點）', '國際標竿與證據', '首案切入', '評等', '建議', '文獻']
CW = [1.85, 3.25, 2.6, 1.95, 1.15, 0.7, 0.95]
y0, rh = 1.4, 0.94
x = 0.45
for j, h in enumerate(HEAD):
    tb(s, x, y0, CW[j] - 0.04, 0.4, h, size=10.5, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE,
       align=PP_ALIGN.CENTER if j >= 4 else PP_ALIGN.LEFT, margin=0.06)
    x += CW[j]
for i, (name, why, bench, land, rate, wave, refs) in enumerate(SVC):
    y = y0 + 0.44 + i * rh; x = 0.45
    first = wave == '第一波'
    bg = LG if i % 2 == 0 else WHITE
    cells = [(name, dict(size=11, bold=True, color=OR if first else INK)),
             (why, dict(size=10)), (bench, dict(size=10)),
             (land, dict(size=10, bold=first)),
             ('  '.join(rate), dict(size=15, color=OR, align=PP_ALIGN.CENTER)),
             (wave, dict(size=10.5, bold=True, color=WHITE if first else OR, align=PP_ALIGN.CENTER, fill=OR if first else TINT)),
             (cites(refs), dict(size=10, align=PP_ALIGN.CENTER))]
    for j, (v, st) in enumerate(cells):
        tb(s, x, y, CW[j] - 0.04, rh - 0.04, v, size=st['size'], bold=st.get('bold', False), color=st.get('color', INK),
           align=st.get('align', PP_ALIGN.LEFT), fill=st.get('fill', bg), anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
        x += CW[j]
tb(s, 0.45, y0 + 0.44 + 5 * rh + 0.02, 12.45, 0.3,
   '評等依序為需求迫切度、與既有服務距離、首案週期：● 有利　◐ 中等　○ 不利（團隊判斷，請會議校準）。文獻編號可點擊，完整清單見附錄 A。', size=9, color=GREY)
# ───────── 第 3 頁 競爭者 ─────────
s = content('MBB 與四大都在推代理型 AI 藍圖；我們以在地確信、IFRS 17 資料與駐點人力區隔')
COMP = [('McKinsey', 'MBB', '核保作業系統 UWOS：45 項能力、6 大領域；代理型 AI 核心現代化（原子化代理能力庫）', '商業與專業險核保、核心重寫', '在台缺在地法規確信與長期駐點人力', [('UWOS', 'mck_uwos'), ('核心現代化', 'mckinsey')]),
        ('BCG', 'MBB', '「AI-First 產險公司」三階段藍圖；亞太保險代理型 AI 白皮書（資源 70% 放在人與變革）', '策略藍圖、組織變革', '偏策略與變革設計，系統落地多需另找夥伴', [('亞太白皮書', 'bcg_ap'), ('產險藍圖', 'bcg_pc')]),
        ('Bain', 'MBB', '亞太保險公司以 AI 重塑理賠與客服案例；保險通路 GenAI 年效益逾 500 億美元（2024）', '理賠、客服、通路', '在台交付規模較小，大型核心導入人力有限', [('案例', 'bain'), ('通路研究', 'bain_dist')]),
        ('Deloitte', '四大', 'Agentic AI Blueprint for Insurance：產險／壽險分類，評估→設計→試點→規模化，預建＋客製代理', '全價值鏈代理型 AI 轉型', '預建代理套到台灣客製化舊核心與繁中條款需調適', [('Blueprint', 'dtt_blueprint'), ('亞太壽險', 'dtt_ap')]),
        ('EY', '四大', 'EY Nexus for Insurance：雲原生平台，快速推出新商品與品牌，串接 650+ 保險科技', '新商品、嵌入式通路、核心延伸', '偏新業務上線，舊核心遷移與 IFRS 17 整合難只靠平台', [('Nexus', 'ey_nexus')]),
        ('KPMG', '四大', 'Intelligent Insurance：AI 價值三階段導入；Trusted AI 十大支柱治理框架', 'AI 導入路徑與治理', '公開資料未見保險專屬代理型產品', [('Intelligent Insurance', 'kpmg_ii'), ('Trusted AI', 'kpmg_trust')]),
        ('PwC（所內）', '四大', 'Agent OS 代理編排平台（跨廠商串接、權限控管、監督儀表板）；IFRS 17／TW-ICS 與 AI 治理服務', '治理＋財務資料＋落地', '可協作：把 Agent OS 與確信能力帶進第一波服務包', [('Agent OS', 'pwc_agentos'), ('IFRS 17', 'pwc_ifrs17')])]
HEAD = ['機構', '保險 AI 主力產品／方法', '主攻', '可能侷限（推論）', '官網／原文']
CW = [1.45, 4.6, 2.05, 2.75, 1.6]
y0, rh = 1.4, 0.66
x = 0.45
for j, h in enumerate(HEAD):
    tb(s, x, y0, CW[j] - 0.04, 0.38, h, size=11, bold=True, color=WHITE, fill=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.08)
    x += CW[j]
for i, (name, grp, prod, focus, gap, links) in enumerate(COMP):
    y = y0 + 0.42 + i * rh; x = 0.45
    ours = name.startswith('PwC')
    bg = TINT if ours else (LG if i % 2 == 0 else WHITE)
    cells = [[[(name, {'bold': True, 'size': 11.5, 'color': OR if ours else INK})], [(grp, {'size': 9, 'color': GREY})]],
             prod, focus, gap,
             [[(lab, {'link': URL[key], 'size': 9.5})] for lab, key in links]]
    for j, v in enumerate(cells):
        tb(s, x, y, CW[j] - 0.04, rh - 0.04, v, size=10, fill=bg, anchor=MSO_ANCHOR.MIDDLE, margin=0.07)
        x += CW[j]
    if i == 2:
        hline(s, 0.45, y + rh - 0.02, 12.9, y + rh - 0.02, color=OR, width=1.5)
tb(s, 0.45, y0 + 0.42 + 7 * rh + 0.04, 12.45, 0.3, '「可能侷限」為依公開資料的推論，僅供內部討論，對外勿點名競爭者。點選右欄可開啟各機構官網或原文。', size=9, color=GREY)

# ───────── 附錄 A 文獻清單 ─────────
s = content('附錄 A｜文獻清單')
half = (len(REFS) + 1) // 2
for i, (cat, title, org, url, ok) in enumerate(REFS):
    col, row = divmod(i, half)
    x = 0.45 + col * 6.3; y = 1.35 + row * 0.47
    tb(s, x, y, 0.45, 0.42, f'[{i + 1}]', size=10, bold=True, color=OR, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    tb(s, x + 0.45, y, 5.75, 0.42, [[(title + ('　※' if not ok else ''), {'link': url, 'size': 10})],
                                   [(f'{cat}｜{org}', {'size': 8.5, 'color': GREY})]], size=10, anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
tb(s, 0.45, 6.62, 12.45, 0.35, '※ 取自原研究連結庫，網路搜尋未能確認確切網址，引用前請點開確認。競爭者官網連結見第 3 頁。', size=9, color=GREY, margin=0.03)

# 同步輸出文獻清單 Markdown
md = ['# 保險科技服務優先順序：文獻清單', '', '主管精簡版第 2 頁文獻編號對應本清單。※ 為網路搜尋未能確認確切網址者。', '',
      '| # | 類別 | 文獻 | 機構與日期 |', '|---|---|---|---|']
md += [f'| {i + 1} | {c} | [{t}]({u}){" ※" if not ok else ""} | {o} |' for i, (c, t, o, u, ok) in enumerate(REFS)]
open(os.path.join(os.path.dirname(os.path.abspath(OUT)), '文獻清單.md'), 'w').write('\n'.join(md) + '\n')

p.save(OUT)
print(OUT, len(p.slides), 'slides')
