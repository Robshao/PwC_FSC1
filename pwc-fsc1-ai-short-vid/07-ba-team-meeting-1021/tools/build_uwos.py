import sys, copy
from pptx import Presentation
from pptx.oxml.ns import qn

src, out = sys.argv[1], sys.argv[2]
p = Presentation(src)

def all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from all_shapes(sh.shapes)

def find(slide, sid):
    for sh in all_shapes(slide.shapes):
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)

def set_text(sh, paras):
    """Replace text, keeping the first paragraph's pPr and first run's rPr."""
    txBody = sh.text_frame._txBody
    ps = txBody.findall(qn('a:p'))
    proto = ps[0]
    for x in ps:
        txBody.remove(x)
    for text in paras:
        np_ = copy.deepcopy(proto)
        runs = np_.findall(qn('a:r'))
        for r in runs[1:]:
            np_.remove(r)
        for tag in ('a:br', 'a:fld'):
            for x in np_.findall(qn(tag)):
                np_.remove(x)
        if runs:
            runs[0].find(qn('a:t')).text = text
        else:
            r = np_.makeelement(qn('a:r'), {})
            t = r.makeelement(qn('a:t'), {}); t.text = text; r.append(t)
            end = np_.find(qn('a:endParaRPr'))
            if end is not None: end.addprevious(r)
            else: np_.append(r)
        txBody.append(np_)

def delete(slide, sid):
    el = find(slide, sid)._element
    el.getparent().remove(el)

def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text

FOOTER = "FSC1 BA 組會｜McKinsey UWOS 文章分享（資料來源：McKinsey & Company, 2026）"

# ---------- Slide 1 (layout from template p.23) ----------
s1 = p.slides[0]
set_text(find(s1, 6), ["McKinsey UWOS：核保轉向「機器優先、人類治理」"])
set_text(find(s1, 9), ["商業險與特種險的核保流程，多半和十年前差不多。McKinsey 認為 Agentic AI 已能執行多步驟工作，保險公司可以把核保重新設計成 UWOS（Underwriting Operating System）：由一個智慧層統籌從受理到出單的流程，人只在需要判斷的地方介入。"])
set_text(find(s1, 17), ["Machine-first, human-governed: AI handles routine work end to end; underwriters focus on portfolio management and complex decisions"])
set_text(find(s1, 19), ["為什麼現在要改"])
set_text(find(s1, 4), ["流程停滯"])
set_text(find(s1, 5), ["營運壓力"])
set_text(find(s1, 7), ["技術轉折"])
set_text(find(s1, 8), ["送件在佇列中排隊", "人工摘事實、追缺漏", "一件要數天到數週"])
set_text(find(s1, 10), ["擴大規模不能加成本", "核保人才短缺", "舊系統、流程零碎"])
set_text(find(s1, 11), ["Agentic AI 跑多步驟", "核對證據、調出前例", "自動準備核保建議"])
set_text(find(s1, 20), ["三條車道：依案件動態分派"])
set_text(find(s1, 14), ["直通車道：簡單、可重複的案件"])
set_text(find(s1, 35), ["在風險胃納、授權與信心門檻內，系統完成受理到出單，人幾乎不介入。"])
set_text(find(s1, 15), ["輔助決策車道：中等複雜度"])
set_text(find(s1, 33), ["系統整理事實、標出不確定處並提出建議，人快速核准、修改或拒絕。"])
set_text(find(s1, 13), ["專家判斷車道：複雜或策略性"])
set_text(find(s1, 36), ["系統備好證據與選項，核保專家主導判斷與談判，時間從行政轉向決策。"])
set_text(find(s1, 21), ["人的角色改變，", "但沒有消失"])
set_text(find(s1, 37), ["核保主管：從處理個案到監督"])
set_text(find(s1, 42), ["形塑業務組合、確保核保品質，並調整風險胃納、定價與經紀人策略。"])
set_text(find(s1, 38), ["資深專家：成為例外處理者"])
set_text(find(s1, 41), ["專注真正需要判斷、談判與經驗的複雜或客製化案件。"])
set_text(find(s1, 39), ["機器：承擔大部分處理工作"])
set_text(find(s1, 40), ["六大領域、45 項能力由智慧層串起，決策更快、更一致，重工更少。"])
set_text(find(s1, 12), [FOOTER])
find(s1, 2).text_frame._txBody.find(qn('a:p')).find(qn('a:fld')).find(qn('a:t')).text = "1"
set_notes(s1, "\n".join([
    "【一個重點】McKinsey 提出 UWOS（核保作業系統）：用 Agentic AI 把核保從逐案人工處理，改成機器優先、人類治理。",
    "為什麼現在：流程十年沒變、成本與人才壓力、Agentic AI 已能跑多步驟流程。",
    "怎麼運作：案件依複雜度、信心程度、所需授權層級分到三條車道，直通、輔助決策、專家判斷。",
    "人的角色：核保主管轉為監督，資深專家處理例外，機器承擔大部分處理。",
    "原文把 UWOS 拆成 6 大領域（受理、風險選擇、定價、報價出單、業務組合管理、營運績效管理）共 45 項能力。",
]))

# ---------- Slide 2 (layout from template p.114) ----------
s2 = p.slides[1]
for sid in (1049, 1051, 1052, 1053, 1056, 1057, 1074, 1082):   # hidden leftovers from the source deck
    delete(s2, sid)
set_text(find(s2, 4), ["從願景到落地：先選一個燈塔案例，再逐步擴大"])
set_text(find(s2, 49), ["McKinsey 建議不要等整個企業資料模型完美才開始，而是從高件量、資料齊全的流程切入，先定好控制機制，再擴大到更多領域。先行動的保險公司，將主導未來十年的速度、服務與核保紀律。"])
set_text(find(s2, 1077), ["一個重點"])
set_text(find(s2, 1076), ["UWOS 不是多加一個工具，而是重新分配人與機器的工作：例行交給系統，判斷留給人。", "原文拆成 6 大領域、45 項能力，由一個智慧層串起。"])
set_text(find(s2, 1078), ["一個啟發"])
set_text(find(s2, 1075), ["AI 導入的關鍵在分派規則與授權門檻：哪些自動通過、哪些交給人，要先講清楚並保留依據，客戶才會信任結果。"])
set_text(find(s2, 1062), ["從願景到落地的四個步驟："])
set_text(find(s2, 8), ["1 選一個", "燈塔案例"])
set_text(find(s2, 13), ["2 建立最低", "限度資料基礎"])
set_text(find(s2, 21), ["3 定義控制、", "授權與覆核"])
set_text(find(s2, 28), ["4 先代理化", "1–2 個流程"])
set_text(find(s2, 50), ["過去：逐案人工處理…"])
set_text(find(s2, 61), ["…將會變成…"])
set_text(find(s2, 51), ["UWOS：機器優先、人類治理"])
set_text(find(s2, 62), ["送件在佇列中排隊等人處理"])
set_text(find(s2, 1026), ["核保人員手動摘資料、追補缺件"])
set_text(find(s2, 1029), ["案件靠人工在各階段之間移動"])
set_text(find(s2, 1042), ["依複雜度與信心程度動態分派"])
set_text(find(s2, 1045), ["例行案件從受理到出單一路處理"])
set_text(find(s2, 1093), ["人專注判斷、談判與業務組合"])
set_text(find(s2, 1094), ["可應用做法：寫需求時，標出每一步是「系統自動」、「AI 輔助＋人確認」或「人判斷」"])
set_text(find(s2, 1095), ["AI 產出要可追溯：保留依據與前例，人工 review 才有基礎"])
set_text(find(s2, 1096), ["先挑一個高件量、資料齊全的流程試做，驗證後再擴大"])
set_text(find(s2, 2), [FOOTER])
find(s2, 1092).text_frame._txBody.find(qn('a:p')).find(qn('a:fld')).find(qn('a:t')).text = "2"
set_notes(s2, "\n".join([
    "落地四步驟：",
    "1. 選燈塔案例：送件頻繁、資料齊全、處理週期或一致性很重要的險種、客群或續保業務。",
    "2. 最低資料基礎：只做第一批流程需要的文件、欄位、規則、定價工具與決策紀錄。",
    "3. 定義控制：信心門檻、授權規則、人工推翻流程，以及 human-in-the-loop 的角色。",
    "4. 啟動目標流程：先代理化 1–2 個高影響力領域（受理、風險選擇或定價），依商業價值、資料準備度與可行性排序。",
    "【一個啟發】AI 輔助能否被客戶接受，取決於分派規則、授權門檻與可追溯的依據，這和我們組會談到的人工 review、建立客戶信任是同一件事。",
    "【一個可應用做法】寫需求或設計流程時，標出每一步由系統自動、AI 輔助＋人確認，或由人判斷。",
]))

p.save(out)
