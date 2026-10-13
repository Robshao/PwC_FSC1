import copy, sys
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
ORANGE = RGBColor(0xFD, 0x51, 0x08); ORANGE2 = RGBColor(0xFE, 0x7C, 0x39); ORANGE3 = RGBColor(0xFF, 0xAA, 0x72)
PEACH = RGBColor(0xFF, 0xE8, 0xD6); GREY = RGBColor(0x5F, 0x66, 0x70); GREY_L = RGBColor(0xCB, 0xD1, 0xD6)
BLACK = RGBColor(0x00, 0x00, 0x00); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
EA = "Microsoft JhengHei"
FOOTER = "AI 短影音製作流程"

prs = Presentation(SRC)

# keep the template cover (slide 1), drop every other sample slide
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[1:]:
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)

def ea(run, latin=None):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea"):
        el = rPr.find(qn(tag))
        if el is not None: rPr.remove(el)
    if latin:
        e = etree.SubElement(rPr, qn("a:latin")); e.set("typeface", latin)
    e = etree.SubElement(rPr, qn("a:ea")); e.set("typeface", EA)

def set_text(tf, paras, size=None, bold=None, color=None, latin=None, align=None):
    """paras: list of str or (str, dict) ; dict keys size,bold,color,level,latin"""
    tf.clear()
    for i, p in enumerate(paras):
        txt, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if "level" in o: para.level = o["level"]
        if align or o.get("align"): para.alignment = o.get("align", align)
        r = para.add_run(); r.text = txt
        f = r.font
        s = o.get("size", size); b = o.get("bold", bold); c = o.get("color", color)
        if s: f.size = Pt(s)
        if b is not None: f.bold = b
        if c is not None: f.color.rgb = c
        ea(r, o.get("latin", latin))

def box(slide, x, y, w, h, paras, fill=None, anchor=MSO_ANCHOR.TOP, shape=MSO_SHAPE.RECTANGLE, line=None, margin=0.15, **kw):
    sh = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"): setattr(tf, f"margin_{side}", Cm(margin))
    set_text(tf, paras, **kw)
    return sh

def footer(slide):
    """clone the layout's footer and slide-number placeholders onto the slide"""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.idx in (3, 4):
            el = copy.deepcopy(ph._element); slide.shapes._spTree.append(el)
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 3:
            set_text(ph.text_frame, [FOOTER])
        if ph.placeholder_format.idx == 4:
            tf = ph.text_frame; tf.clear(); p = tf.paragraphs[0]
            fld = etree.SubElement(p._p, qn("a:fld")); fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"); fld.set("type", "slidenum")
            etree.SubElement(fld, qn("a:t")).text = str(len(prs.slides))

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def connector(slide, x1, y1, x2, y2, color, width=2, dash=None, head=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", dash)
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c

def layout(name):
    for l in prs.slide_layouts:
        if l.name.strip() == name: return l
    raise KeyError(name)

# ---------- 1. cover (template slide, text replaced) ----------
cover = prs.slides[0]
for sh in cover.placeholders:
    if str(sh.placeholder_format.type).split(".")[-1].split(" ")[0] in ("CENTER_TITLE", "TITLE"):
        set_text(sh.text_frame, ["AI 短影音製作流程"])
    elif sh.has_text_frame:
        set_text(sh.text_frame, [("從主題到成片：7 個步驟｜首支約 2.5 個工作天", {"bold": True}),
                                 "PwC FSC1 內部分享｜2026 年 10 月"])
notes(cover, "這份簡報整理我們做客戶分層樣片的完整流程：流程可以標準化成 7 個步驟；第一支約 2.5 個工作天，之後沿用角色、場景與後製模板可縮短到 1–1.5 天。")

# ---------- 2. process map ----------
s = prs.slides.add_slide(layout("只有標題"))
set_text(s.shapes.title.text_frame, ["製作流程圖：7 個步驟"])
box(s, 1.0, 3.3, 31.8, 0.9, ["每一步都有固定的工具與產出；前一步的產出，就是下一步的輸入"], size=14, color=GREY, margin=0, align=PP_ALIGN.LEFT)
X0, W = 1.0, 31.8 / 7
phases = [(0, 3, "前製｜文字工作"), (3, 2, "製作｜AI 生成"), (5, 2, "後製與審查")]
for start, n, label in phases:
    box(s, X0 + start * W + 0.05, 4.6, n * W - 0.1, 1.0, [label], fill=PEACH, anchor=MSO_ANCHOR.MIDDLE,
        size=14, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
steps = [("主題與目的", "主管需求", "目的・觀眾・訊息"), ("大綱・人物誌", "人物誌＋分層框架", "大綱五要素"),
         ("劇本・分鏡", "Claude", "分鏡表＋全套提示詞"), ("參考圖生成", "Gemini", "角色設定＋關鍵畫面"),
         ("影片生成", "Seedance", "4–15 秒有聲片段"), ("後製合成", "ffmpeg／CapCut", "字幕・轉場・配樂"),
         ("審查・迭代", "導演式看片", "修改清單→成片")]
circ_fill = [ORANGE3, ORANGE3, ORANGE3, ORANGE2, ORANGE2, ORANGE, ORANGE]
num_col = [BLACK, BLACK, BLACK, WHITE, WHITE, WHITE, WHITE]
LINE_Y, D = 7.7, 1.9
connector(s, X0 + W / 2, LINE_Y, X0 + 6.5 * W, LINE_Y, GREY_L, width=3)
for i, (name, tool, out) in enumerate(steps):
    cx = X0 + W * (i + 0.5)
    box(s, cx - D / 2, LINE_Y - D / 2, D, D, [str(i + 1)], fill=circ_fill[i], shape=MSO_SHAPE.OVAL,
        anchor=MSO_ANCHOR.MIDDLE, size=24, bold=True, color=num_col[i], latin="Georgia", align=PP_ALIGN.CENTER, margin=0)
    box(s, cx - W / 2 + 0.1, 9.0, W - 0.2, 3.2,
        [(name, {"size": 16, "bold": True}), (tool, {"size": 13, "color": GREY}), (out, {"size": 13})],
        align=PP_ALIGN.CENTER, color=BLACK, margin=0.05)
# review loop 7 -> 5
c7, c5, ly = X0 + 6.5 * W, X0 + 4.5 * W, 13.1
connector(s, c7, 12.2, c7, ly, ORANGE, width=2, dash="dash")
connector(s, c7, ly, c5, ly, ORANGE, width=2, dash="dash")
connector(s, c5, ly, c5, 12.2, ORANGE, width=2, dash="dash", head=True)
box(s, c5 - 0.5, 13.35, c7 - c5 + 1.0, 1.6, ["審查不通過 → 回到第 5 步重新生成", "（本片共迭代 11 版）"],
    size=13, bold=True, color=ORANGE, align=PP_ALIGN.CENTER, margin=0)
box(s, 1.0, 13.0, 18.5, 3.2, [("分工原則", {"size": 14, "bold": True, "color": ORANGE}),
    ("Gemini 只做參考圖，Seedance 只做影片；字幕、金融數字與品牌元素一律在後製加上。", {"size": 16})],
    fill=PEACH, color=BLACK, margin=0.45, align=PP_ALIGN.LEFT)
footer(s)
notes(s, "由左到右：前三步是文字工作，決定故事好不好；第四、五步交給 AI 生成畫面與影片；第六、七步是後製與審查。橘色虛線是修改回圈，這支片共跑了 11 個版本。")

# ---------- 3. time estimate ----------
s = prs.slides.add_slide(layout("只有標題"))
set_text(s.shapes.title.text_frame, ["時間估算：一支 45 秒影片"])
rows = [("步驟", "首支", "第二支起", "時間主要花在"),
        ("1 主題與目的", "0.5 小時", "0.5 小時", "確認目的、觀眾、核心訊息"),
        ("2 大綱・人物誌", "1 小時", "0.5 小時", "對齊人物誌與分層框架"),
        ("3 劇本・分鏡", "3 小時", "1.5 小時", "台詞字數配合鏡頭秒數"),
        ("4 參考圖生成", "2 小時", "1 小時", "角色一致性、挑圖"),
        ("5 影片生成", "4 小時", "2 小時", "無聲、手部畸形的重新生成"),
        ("6 後製合成", "4 小時", "2 小時", "字幕對齊、轉場、配樂、螢幕替換"),
        ("7 審查・迭代", "4 小時", "2 小時", "每輪看片→修改清單"),
        ("合計", "18.5 小時", "9.5 小時", "約省一半")]
widths = [5.2, 3.0, 3.2, 9.6]
tb = s.shapes.add_table(len(rows), 4, Cm(1.0), Cm(4.0), Cm(sum(widths)), Cm(0.95 * len(rows))).table
for j, w in enumerate(widths): tb.columns[j].width = Cm(w)
for i, row in enumerate(rows):
    tb.rows[i].height = Cm(0.95)
    for j, val in enumerate(row):
        cell = tb.cell(i, j); cell.margin_left = cell.margin_right = Cm(0.2)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        head, total = i == 0, i == len(rows) - 1
        cell.fill.solid(); cell.fill.fore_color.rgb = ORANGE if head else (PEACH if total else WHITE)
        set_text(cell.text_frame, [val], size=11, bold=head or total, color=WHITE if head else BLACK,
                 align=PP_ALIGN.RIGHT if j in (1, 2) else PP_ALIGN.LEFT)
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT"):
            ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", "0"); etree.SubElement(ln, qn("a:noFill"))
        ln = etree.SubElement(tcPr, qn("a:lnB")); ln.set("w", "6350")
        sf = etree.SubElement(ln, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", "CBD1D6")
        # schema order: lnL lnR lnT lnB must precede fill
        fill = tcPr.find(qn("a:solidFill"))
        if fill is not None: tcPr.remove(fill); tcPr.append(fill)
tblPr = tb._tbl.tblPr
for k in ("firstRow", "bandRow"): tblPr.set(k, "0")
box(s, 23.2, 4.0, 9.6, 3.6, [("首支（從零開始）", {"size": 12}), ("約 2.5 天", {"size": 40, "latin": "Georgia"})],
    fill=ORANGE, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, margin=0.5)
box(s, 23.2, 7.9, 9.6, 3.6, [("第二支起（沿用資產）", {"size": 12, "color": BLACK}), ("1–1.5 天", {"size": 40, "latin": "Georgia", "color": ORANGE})],
    fill=PEACH, anchor=MSO_ANCHOR.MIDDLE, margin=0.5)
box(s, 23.2, 11.8, 9.6, 1.8, ["沿用：顧問角色、客廳與咖啡店場景、字卡與轉場模板、配樂與混音設定"], size=11, color=GREY, margin=0)
box(s, 1.0, 13.7, 31.8, 0.9, ["依本片實際製作經驗估算（10/12–10/13，11 個版本），以一天 8 小時計；不含主管與法遵審核等待、配樂授權採購。"],
    size=9, color=GREY, margin=0)
footer(s)
notes(s, "這是估計值，不是精確工時紀錄。最花時間的是第五到第七步：AI 生成常要重來，後製的字幕對齊和轉場也要來回調。第二支起沿用角色、場景與後製模板，約可省一半。")

# ---------- 4. lessons (Six Content Subtitle) ----------
s = prs.slides.add_slide(layout("Six Content Subtitle"))
lessons = [("角色每個鏡頭都不一樣", "先做角色設定圖；每段提示詞都上傳它，並寫明「臉、髮型、服裝不變」。"),
           ("生成的影片沒有聲音", "提示詞開頭與聲音段都寫「必須有聲音」；下載前先在預覽裡聽過。"),
           ("台詞太長被切掉", "以每秒 4–5 個字估算；一句台詞要放得進 4–7 秒的鏡頭。"),
           ("場景跳接看不懂", "過渡字卡＋掛斷音效＋環境音先進場（J-cut）；倒帶要慢到看得出來。"),
           ("AI 生成的字是亂碼", "金融數字與介面文字一律後製貼上，例如平板螢幕替換。"),
           ("手部畸形、多手指", "改用乾淨的段落；或局部重繪後當首幀再生成一次。")]
content_idx = [10, 11, 12, 13, 14, 16]
for ph in s.placeholders:
    idx = ph.placeholder_format.idx
    if idx == 0: set_text(ph.text_frame, ["常見問題與標準解法"])
    elif idx == 15: set_text(ph.text_frame, ["這支片實際遇到的 6 個問題：前三個在生成階段預防，後三個在後製處理"])
    elif idx in content_idx:
        h, b = lessons[content_idx.index(idx)]
        set_text(ph.text_frame, [(h, {"bold": True, "level": 0, "size": 15}), (b, {"level": 1, "size": 13})])
for i in range(6):
    box(s, 1.0 + i * 5.36, 4.75, 5.1, 0.12, [""], fill=ORANGE3 if i < 3 else ORANGE, margin=0)
footer(s)
notes(s, "這六個都是這支片真的遇到的問題。前三個寫進提示詞範本，在生成階段預防；後三個是後製的標準做法。")

# ---------- 5. next video (Agenda 2) ----------
s = prs.slides.add_slide(layout("Agenda 2"))
for ph in s.placeholders:
    idx = ph.placeholder_format.idx
    if idx == 0: set_text(ph.text_frame, ["下一支影片：先準備 5 項輸入"])
    elif idx == 12:
        items = [("主題與目的", "要說服誰？看完要讓他們做什麼決定？"),
                 ("目標客群與人物誌", "例：新婚首購族、退休準備族"),
                 ("片長與使用場景", "提案開場、研討會，或社群短影音"),
                 ("品牌與法遵限制", "標誌露出、疾病情節、數字與配樂授權"),
                 ("可沿用的資產", "顧問角色、客廳與咖啡店場景、字卡與轉場模板")]
        tf = ph.text_frame; tf.clear()
        for k, (h, d) in enumerate(items):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.space_after = Pt(10)
            for txt, sz, b, col, lat in ((f"{k+1}  ", 20, True, ORANGE, "Georgia"), (h + "｜", 16, True, BLACK, None), (d, 14, False, GREY, None)):
                r = para.add_run(); r.text = txt; r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = col; ea(r, lat)
box(s, 14.4, 13.2, 17.4, 2.0, [("5 項到齊後，從第 3 步「劇本・分鏡」開始，約 1–1.5 個工作天交付初版", {"size": 14, "bold": True})],
    fill=WHITE, color=BLACK, anchor=MSO_ANCHOR.MIDDLE, margin=0.45)
footer(s)
notes(s, "下一支片只要先確認這五項，就能直接從劇本和分鏡開始。第五項最能省時間：已經做好的顧問角色和場景都可以直接沿用。")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
